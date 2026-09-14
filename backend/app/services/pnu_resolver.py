"""PNU 확정 파이프라인 (표준데이터셋 제작·갱신 프로세스 설계서 §7.2 「PNU 기준」).

박진주 대표 설계서 §7.2 의 절차를 그대로 구현한다. 이 절차는 룰북 §3 이 금지하는
「문자열 유사매칭 확정」이 아니라, §3 이 요구하는 PNU 우선키를 **결정적 조립 +
공간 조인**으로 실제로 만들어내는 재현 가능한 절차다.

절차(설계서 §7.2)
------------------
① 입력 PNU 검증      원천에 PNU 가 있으면 19자리 형식 확인 + 연속지적도 실재 검증.
② 지번주소 → PNU 조립  PNU(19) = 법정동코드(10) + 대지구분(1) + 본번(4) + 부번(4).
                       대지구분: 일반 1, 산 2. 본번·부번 앞자리 0 채워 4자리.
                       법정동명은 최장 일치로 찾는다(외부 조회 없이 재현 가능).
③ 좌표 확보          원천 좌표가 유효하면 사용, 없으면 지오코딩으로 보완.
④ 연속지적도 대조     좌표를 EPSG:5186 으로 변환해 필지와 대조.
                       **좌표를 포함하는 필지가 단 하나일 때만** 유효한 대조로 인정.
⑤ 교차 검증          조립 PNU 와 대조 필지가 일치하면 확정.
                       불일치·필지 미특정이면 확정하지 않고 검토대상으로 분리.
                       조립 PNU 를 못 구했어도 좌표가 단일 필지에 포함되면 그 필지
                       PNU 로 확정. 좌표가 없어도 조립 PNU 가 실재하면 내부 대표점을
                       좌표로 파생(CADASTRAL_REPRESENTATIVE_POINT).
⑥ 실패 사유 구분      후속 조치가 달라지므로 원인을 나눠 기록한다.

절대 원칙
- 후보 필지가 2개 이상이면 확정하지 않는다(검토대상). 경계 공유 필지에서 잘못
  확정하는 것을 막는다. 이 판정은 CadastralLocalStore.sole_parcel_at() 이 「단
  하나일 때만」 반환하도록 강제한다(parcel_at() 의 최소면적 우선과 다르다).
- 조립 PNU 와 공간조인 결과가 불일치하면 확정하지 않는다.
- 지오코딩·좌표 결측을 빈 결과(=확정)로 삼키지 않는다. 확정하지 못한 건은 삭제하지
  않고 사유와 함께 검토대상/실패로 격리한다(룰북 §7 ⑤).

이 모듈은 외부 네트워크를 호출하지 않는다. 좌표가 없어도 조립 PNU 가 연속지적도에
실재하면 내부 대표점으로 확정하므로, 운영 오프라인 전제(설계서 §1.4)를 지킨다.
"""

from __future__ import annotations

import os
import re
from collections.abc import Iterable, Mapping
from dataclasses import dataclass, field
from pathlib import Path
from typing import Literal, Protocol

from app.models import Coordinates

# 지연 import 로 순환/무거운 의존을 피한다(openpyxl·shapely 는 실제 사용 시 로드).


# ── 상수 ─────────────────────────────────────────────────────────────────────

ENV_LEGAL_DONG = "LH_LEGAL_DONG_PATH"

# 좌표 유효범위(설계서 §6.3). 벗어나면 좌표계 오인이라 좌표 없음으로 취급한다.
KOREA_LAT_MIN, KOREA_LAT_MAX = 33.0, 39.5
KOREA_LNG_MIN, KOREA_LNG_MAX = 124.0, 132.0

# 확정 알고리즘 버전. resolve() 의 확정/검토 분기 규칙이 바뀌면 올린다. PNU 확정
# 결과 사이드카 캐시 키에 포함돼, 코드가 바뀌면 낡은 캐시가 자동 무효화된다.
# v2: 좌표가 있으나 공간조인 0건이면 대표점 확정 대신 검토대상으로 분리(모순 좌표
# 확정 방지).
# v3: 주소 → PNU 조립에 API 경로(address_pnu)를 먼저 쓰는 분기를 넣었다. 표가 없던
# 환경에서 조립이 성립하게 되므로 이전 캐시(전건 실패)를 재사용하면 안 된다.
RESOLVER_ALGORITHM_VERSION = 3

# 확정 경로 코드(판정 근거로 화면에 나간다).
METHOD_INPUT_PNU = "INPUT_PNU"  # 원천 PNU 검증 통과
METHOD_ADDRESS_ASSEMBLY = "ADDRESS_ASSEMBLY"  # 조립 PNU ↔ 공간조인 일치
METHOD_COORDINATE_SPATIAL_JOIN = "COORDINATE_SPATIAL_JOIN"  # 좌표 단일 필지 포함
METHOD_CADASTRAL_REPRESENTATIVE_POINT = "CADASTRAL_REPRESENTATIVE_POINT"  # 대표점 파생

# 검토대상 사유(확정하지 않되 실패도 아님).
REVIEW_MULTIPLE_PARCELS = "MULTIPLE_PARCELS"  # 후보 필지 2개 이상
REVIEW_ASSEMBLY_SPATIAL_MISMATCH = "ASSEMBLY_SPATIAL_MISMATCH"  # 조립 ↔ 공간조인 불일치
REVIEW_INPUT_PNU_MISMATCH = "INPUT_PNU_MISMATCH"  # 입력 PNU ↔ 공간조인 불일치
# 원천 좌표가 있는데 어떤 필지에도 안 들어감(공간조인 0건). 좌표가 대표점과
# 모순되므로 확정하지 않는다. 좌표가 없는 레코드만 대표점 확정을 허용한다.
REVIEW_COORDINATE_OUTSIDE_PARCEL = "COORDINATE_OUTSIDE_PARCEL"

# 실패 사유 3종(설계서 §7.2 ⑥ 표 그대로).
FAIL_JIBUN_NOT_IN_CADASTRAL = "JIBUN_NOT_IN_CADASTRAL"
# 지번은 있으나 해당 필지가 지적도에 없음(같은 본번 다른 부번 존재) → 지번 변경·합병·분할
FAIL_BONBUN_NOT_IN_CADASTRAL = "BONBUN_NOT_IN_CADASTRAL"
# 해당 본번 자체가 지적도에 없음 → 주소 지번 오류·폐쇄지번
FAIL_NO_JIBUN = "NO_JIBUN"
# 지번 없는 주소(도로명·산업단지 블록 표기 등) → 조립 불가

FAILURE_ACTIONS: dict[str, str] = {
    FAIL_JIBUN_NOT_IN_CADASTRAL: "지번 변경 이력 확인",
    FAIL_BONBUN_NOT_IN_CADASTRAL: "원천 주소 재확인",
    FAIL_NO_JIBUN: "지번주소 확보",
}

FAILURE_LABELS: dict[str, str] = {
    FAIL_JIBUN_NOT_IN_CADASTRAL: "지번은 있으나 해당 필지가 지적도에 없음(같은 본번의 다른 부번은 존재)",
    FAIL_BONBUN_NOT_IN_CADASTRAL: "해당 본번 자체가 지적도에 없음",
    FAIL_NO_JIBUN: "지번이 없는 주소(도로명·산업단지 블록 표기 등) — 조립 불가",
}

ResolveStatus = Literal["confirmed", "review", "failed"]


# ── 주소 정규화 (설계서 §6.2) ────────────────────────────────────────────────


def normalize_address(address: str) -> str:
    """행정구역 개편 표기를 통일한다(설계서 §6.2). 원문 의미는 보존한다."""

    text = (address or "").strip()
    if not text:
        return ""
    text = text.replace("전라북도", "전북특별자치도")
    # "전북 " → "전북특별자치도 " (이미 정식표기면 건드리지 않는다)
    if text.startswith("전북 "):
        text = "전북특별자치도 " + text[len("전북 "):]
    text = text.replace("전라남도", "전남").replace("충청남도", "충남")
    return text


def _squash(text: str) -> str:
    """공백을 모두 제거한다(법정동 최장 일치·지번 파싱용)."""

    return re.sub(r"\s+", "", text or "")


# ── 법정동 코드 색인 (지번주소 → 10자리 코드, 최장 일치) ──────────────────────


@dataclass(frozen=True)
class LegalDongMatch:
    code: str  # 10자리 법정동코드
    name: str  # 매칭된 법정동 전체명(공백 제거)
    end: int  # 정규화 주소에서 매칭 끝 위치(지번 파싱 시작점)


class LegalDongIndex:
    """법정동명 → 10자리 코드 색인. 최장 일치로 코드를 확정한다(설계서 §7.2 ②).

    행정안전부 법정동코드표(시도·시군구·읍면동·리 분리 컬럼)를 읽어, 각 읍면동/리
    레코드에 대해 (시군구+읍면동+리) 를 공백 제거한 키를 만든다. 조회는 먼저
    시군구가 주소에 있는지 보고, 그 시군구의 읍면동/리 중 가장 긴 키가 주소에
    나타나는 것을 고른다. 외부 조회 없이 결정되므로 재현 가능하다.
    """

    def __init__(self, entries: Iterable[tuple[str, str, str, str, str]]):
        # entries: (code, sido, sgg, emd, ri)
        self._sgg_dongs: dict[str, list[tuple[str, str]]] = {}
        self._sgg_list: list[str] = []
        seen_sgg: set[str] = set()
        for code, _sido, sgg, emd, ri in entries:
            code = (code or "").strip()
            if len(code) != 10 or not code.isdigit():
                continue
            if not emd:  # 읍면동/리 레벨만(시도·시군구 레벨은 지번 조립 대상 아님)
                continue
            sgg_key = _squash(sgg or "")
            if not sgg_key:
                continue
            dong_key = _squash((sgg or "") + (emd or "") + (ri or ""))
            self._sgg_dongs.setdefault(sgg_key, []).append((dong_key, code))
            if sgg_key not in seen_sgg:
                seen_sgg.add(sgg_key)
                self._sgg_list.append(sgg_key)
        # 최장 일치를 위해 시군구·읍면동 키를 길이 내림차순으로 정렬한다.
        self._sgg_list.sort(key=len, reverse=True)
        for sgg_key in self._sgg_dongs:
            # 중복 키는 첫 코드만 남긴다(정규화 결과가 같으면 동일 법정동).
            deduped: dict[str, str] = {}
            for dong_key, code in self._sgg_dongs[sgg_key]:
                deduped.setdefault(dong_key, code)
            self._sgg_dongs[sgg_key] = sorted(
                deduped.items(), key=lambda kv: len(kv[0]), reverse=True
            )

    @property
    def size(self) -> int:
        return sum(len(v) for v in self._sgg_dongs.values())

    def match(self, address: str) -> LegalDongMatch | None:
        """정규화·공백제거 주소에서 최장 일치 법정동을 찾는다."""

        squashed = _squash(normalize_address(address))
        if not squashed:
            return None
        for sgg_key in self._sgg_list:
            if sgg_key not in squashed:
                continue
            for dong_key, code in self._sgg_dongs[sgg_key]:
                pos = squashed.find(dong_key)
                if pos >= 0:
                    return LegalDongMatch(
                        code=code, name=dong_key, end=pos + len(dong_key)
                    )
        return None

    @classmethod
    def from_xlsx(cls, path: str | os.PathLike[str]) -> "LegalDongIndex | None":
        """법정동코드 xlsx 를 읽어 색인을 만든다. 파일이 없으면 None."""

        p = Path(path)
        if not p.is_file():
            return None
        from openpyxl import load_workbook

        try:
            wb = load_workbook(p, read_only=True, data_only=True)
        except Exception:
            return None
        try:
            ws = wb[wb.sheetnames[0]]
            it = ws.iter_rows(values_only=True)
            next(it, None)  # 헤더: 법정동코드·시도명·시군구명·읍면동명·리명·순번·생성일자
            rows: list[tuple[str, str, str, str, str]] = []
            for row in it:
                if not row:
                    continue
                code = str(row[0]).strip() if row[0] is not None else ""
                sido = str(row[1]).strip() if len(row) > 1 and row[1] else ""
                sgg = str(row[2]).strip() if len(row) > 2 and row[2] else ""
                emd = str(row[3]).strip() if len(row) > 3 and row[3] else ""
                ri = str(row[4]).strip() if len(row) > 4 and row[4] else ""
                rows.append((code, sido, sgg, emd, ri))
            return cls(rows)
        finally:
            wb.close()

    @classmethod
    def from_env(
        cls, env: Mapping[str, str] | None = None
    ) -> "LegalDongIndex | None":
        source = os.environ if env is None else env
        path = source.get(ENV_LEGAL_DONG)
        if not path:
            return None
        return cls.from_xlsx(path)


# ── 지번 파싱 / PNU 조립 ─────────────────────────────────────────────────────

# 법정동명 뒤에 오는 지번: 선택적 '산' + 본번(-부번). 번지·지목·건물명은 무시한다.
_JIBUN_RE = re.compile(r"^(산)?\s*(\d+)(?:\s*-\s*(\d+))?")


@dataclass(frozen=True)
class AssembledPnu:
    pnu: str
    dong_code: str
    is_mountain: bool
    bonbun: int
    bubun: int

    @property
    def bonbun_prefix(self) -> str:
        """법정동(10)+대지구분(1)+본번(4). 같은 본번 다른 부번 존재 여부 조회용."""

        daeji = "2" if self.is_mountain else "1"
        return f"{self.dong_code}{daeji}{self.bonbun:04d}"


def assemble_pnu(address: str, legal_dong: LegalDongIndex) -> AssembledPnu | None:
    """지번주소로 19자리 PNU 를 조립한다(설계서 §7.2 ②). 조립 불가면 None."""

    match = legal_dong.match(address)
    if match is None:
        return None
    squashed = _squash(normalize_address(address))
    remainder = squashed[match.end:]
    m = _JIBUN_RE.match(remainder)
    if not m:
        return None
    is_mountain = m.group(1) is not None
    bonbun = int(m.group(2))
    bubun = int(m.group(3) or 0)
    if bonbun == 0:
        return None
    daeji = "2" if is_mountain else "1"
    pnu = f"{match.code}{daeji}{bonbun:04d}{bubun:04d}"
    return AssembledPnu(
        pnu=pnu,
        dong_code=match.code,
        is_mountain=is_mountain,
        bonbun=bonbun,
        bubun=bubun,
    )


def validate_pnu(pnu: str) -> str:
    """원천 PNU 형식 검증(설계서 §7.2 ①). 19자리 숫자면 그대로, 아니면 빈 문자열."""

    value = (pnu or "").strip()
    return value if len(value) == 19 and value.isdigit() else ""


# ── 지적도 백엔드 프로토콜 ───────────────────────────────────────────────────


class CadastralBackend(Protocol):
    """PNU 확정에 필요한 연속지적도 조회 인터페이스.

    CadastralLocalStore 가 이 프로토콜을 만족한다. 테스트는 가짜 백엔드를 주입한다.
    """

    def parcel_by_pnu(self, pnu: str):  # -> LocalParcel | None
        ...

    def sole_parcel_at(self, lat: float, lng: float):  # -> tuple[LocalParcel|None, int]
        ...

    def representative_point(self, parcel) -> Coordinates | None:
        ...

    def pnu_prefix_exists(self, prefix: str) -> bool:
        ...


# ── 확정 결과 ────────────────────────────────────────────────────────────────


@dataclass(frozen=True)
class PnuResolution:
    """PNU 확정 결과 하나. 확정 경로·실패 사유를 판정 근거로 남긴다."""

    status: ResolveStatus
    pnu: str = ""
    coordinates: Coordinates | None = None
    method: str = ""  # 확정 경로 코드(status=='confirmed' 일 때)
    reason: str = ""  # 검토/실패 사유 코드
    assembled_pnu: str = ""  # 조립 PNU(있으면). 감사·대조용
    note: str = ""

    @property
    def confirmed(self) -> bool:
        return self.status == "confirmed"

    @property
    def action(self) -> str:
        """실패 사유에 따른 후속 조치 문구(설계서 §7.2 ⑥)."""

        return FAILURE_ACTIONS.get(self.reason, "")


# ── 해석기 ───────────────────────────────────────────────────────────────────


class Geocoder(Protocol):
    def __call__(self, address: str) -> Coordinates | None:
        ...


def _valid_coords(coords: Coordinates | None) -> Coordinates | None:
    if coords is None:
        return None
    if not (KOREA_LAT_MIN <= coords.lat <= KOREA_LAT_MAX):
        return None
    if not (KOREA_LNG_MIN <= coords.lng <= KOREA_LNG_MAX):
        return None
    return coords


class AddressPnuAssembler(Protocol):
    """주소 하나를 PNU 로 조립하는 외부 경로(법정동코드표 대체).

    `LegalDongIndex` 는 법정동코드표 파일이 있어야 하고 문자열 최장 일치로 동을
    고른다. 이 주입점은 주소별로 법정동코드를 직접 돌려주는 원천(카카오 주소검색
    API 등)을 쓰기 위한 자리다. 표가 없어도 조립이 성립한다.
    """

    def __call__(self, address: str) -> "AssembledPnu | None":
        ...


class PnuResolver:
    """설계서 §7.2 절차를 실행한다. 지적도 백엔드가 없으면 확정하지 않는다.

    주소 → PNU 조립 경로는 둘 중 하나면 된다.
      · `legal_dong`  — 법정동코드표 파일(오프라인·결정적)
      · `address_pnu` — 주소별 법정동코드를 주는 API(표 없이 동작)
    둘 다 주면 `address_pnu` 를 먼저 쓰고, 실패하면 표로 물러선다.
    """

    def __init__(
        self,
        legal_dong: LegalDongIndex | None,
        cadastral: CadastralBackend | None,
        geocoder: Geocoder | None = None,
        address_pnu: AddressPnuAssembler | None = None,
    ) -> None:
        self.legal_dong = legal_dong
        self.cadastral = cadastral
        self.geocoder = geocoder
        self.address_pnu = address_pnu

    @property
    def ready(self) -> bool:
        """확정을 시도할 수 있는 상태인지(지적도 백엔드 필수)."""

        return self.cadastral is not None

    # -- 저수준 지적도 접근(백엔드 부재·오류를 삼키지 않고 방어) --

    def _fetch(self, pnu: str):
        """PNU 로 필지를 한 번만 조회한다(대표점 파생 시 재조회를 피한다)."""

        if not pnu or self.cadastral is None:
            return None
        return self.cadastral.parcel_by_pnu(pnu)

    # -- 메인 --

    def resolve(
        self,
        *,
        address: str = "",
        pnu: str = "",
        coordinates: Coordinates | None = None,
    ) -> PnuResolution:
        if self.cadastral is None:
            # 지적도 백엔드가 없으면 실재 검증·공간조인이 불가능하다. 확정하지 않고
            # 실패로 둔다(빈 결과로 삼키지 않는다).
            return PnuResolution(
                status="failed",
                reason=FAIL_NO_JIBUN,
                note="연속지적도 백엔드 미적재 — PNU 확정 불가",
            )

        # ① 입력 PNU 검증
        input_pnu = validate_pnu(pnu)

        # ② 지번주소 → PNU 조립
        # API 경로(주소별 법정동코드)를 먼저 쓰고, 실패하면 법정동코드표로 물러선다.
        # 어느 경로로 조립됐든 이후 검증(③~⑥)은 완전히 같다.
        assembled: AssembledPnu | None = None
        if self.address_pnu is not None and address:
            assembled = self.address_pnu(address)
        if assembled is None and self.legal_dong is not None:
            assembled = assemble_pnu(address, self.legal_dong)
        assembled_pnu = assembled.pnu if assembled else ""
        assembled_parcel = self._fetch(assembled_pnu)
        assembled_real = assembled_parcel is not None

        # ③ 좌표 확보(원천 좌표 우선, 없으면 지오코딩 보완)
        coords = _valid_coords(coordinates)
        if coords is None and self.geocoder is not None and address:
            coords = _valid_coords(self.geocoder(address))

        # ④ 연속지적도 대조(단 하나일 때만 유효)
        join_pnu = ""
        join_count = 0
        join_coords: Coordinates | None = None
        if coords is not None:
            parcel, join_count = self.cadastral.sole_parcel_at(coords.lat, coords.lng)
            if join_count == 1 and parcel is not None:
                join_pnu = parcel.pnu
                join_coords = coords

        # ⑤ 교차 검증 --------------------------------------------------------

        # 입력 PNU 가 실재하면 그것이 정본이다(설계서 §7.2 ①).
        input_parcel = self._fetch(input_pnu) if input_pnu else None
        if input_pnu and input_parcel is not None:
            # 좌표가 단일 필지에 있고 입력 PNU 와 다르면 불일치 검토대상으로 뺀다.
            if join_pnu and join_pnu != input_pnu:
                return PnuResolution(
                    status="review",
                    reason=REVIEW_INPUT_PNU_MISMATCH,
                    pnu="",
                    assembled_pnu=assembled_pnu,
                    note=(
                        f"입력 PNU({input_pnu})와 좌표 공간조인 필지({join_pnu})가 "
                        "불일치. 확정하지 않고 검토대상으로 분리."
                    ),
                )
            out_coords = coords or self._rep_of(input_parcel)
            return PnuResolution(
                status="confirmed",
                pnu=input_pnu,
                coordinates=out_coords,
                method=METHOD_INPUT_PNU,
                assembled_pnu=assembled_pnu,
                note="원천 PNU 형식·실재 검증 통과",
            )

        # 좌표가 여러 필지에 걸리면 확정하지 않는다(경계 공유 필지 오확정 방지).
        if coords is not None and join_count > 1:
            return PnuResolution(
                status="review",
                reason=REVIEW_MULTIPLE_PARCELS,
                assembled_pnu=assembled_pnu,
                note=(
                    f"좌표가 {join_count}개 필지에 걸쳐 단일 필지로 특정 불가. "
                    "확정하지 않고 검토대상으로 분리."
                ),
            )

        # 조립 PNU 실재 + 좌표 단일 필지 → 일치/불일치 교차검증
        if assembled_real and join_pnu:
            if assembled_pnu == join_pnu:
                return PnuResolution(
                    status="confirmed",
                    pnu=assembled_pnu,
                    coordinates=join_coords,
                    method=METHOD_ADDRESS_ASSEMBLY,
                    assembled_pnu=assembled_pnu,
                    note="조립 PNU 와 좌표 공간조인 필지 일치",
                )
            return PnuResolution(
                status="review",
                reason=REVIEW_ASSEMBLY_SPATIAL_MISMATCH,
                assembled_pnu=assembled_pnu,
                note=(
                    f"조립 PNU({assembled_pnu})와 좌표 공간조인 필지({join_pnu})가 "
                    "불일치. 확정하지 않고 검토대상으로 분리."
                ),
            )

        # 조립 PNU 는 없지만 좌표가 단일 필지에 포함 → 그 필지로 확정
        if join_pnu and not assembled_real:
            return PnuResolution(
                status="confirmed",
                pnu=join_pnu,
                coordinates=join_coords,
                method=METHOD_COORDINATE_SPATIAL_JOIN,
                assembled_pnu=assembled_pnu,
                note="좌표가 단일 필지에 포함되어 해당 필지 PNU 로 확정",
            )

        # 좌표가 없어도 조립 PNU 가 실재 → 내부 대표점 파생 확정
        if assembled_real and coords is None:
            rep = self._rep_of(assembled_parcel)
            if rep is not None:
                return PnuResolution(
                    status="confirmed",
                    pnu=assembled_pnu,
                    coordinates=rep,
                    method=METHOD_CADASTRAL_REPRESENTATIVE_POINT,
                    assembled_pnu=assembled_pnu,
                    note="조립 PNU 가 지적도에 실재하여 내부 대표점을 좌표로 파생",
                )

        # 좌표가 있으나 어떤 필지에도 안 들어가고(조립 PNU 실재), 공간조인 0건:
        # 원천이 준 좌표가 조립 PNU 대표점과 모순된다. 주소 오타가 우연히 실재
        # 필지로 조립되면 엉뚱한 필지가 조용히 붙을 수 있으므로, 쓸 수 있는 좌표가
        # 있었는데 단일 필지로 특정되지 않은 이 경우는 확정하지 않고 검토대상으로
        # 분리한다(대표점 확정은 애초에 좌표가 없는 레코드에만 허용한다 — 위 블록).
        if assembled_real and coords is not None and join_count == 0:
            return PnuResolution(
                status="review",
                reason=REVIEW_COORDINATE_OUTSIDE_PARCEL,
                assembled_pnu=assembled_pnu,
                note=(
                    f"원천 좌표가 어떤 필지에도 포함되지 않아(공간조인 0건) 조립 "
                    f"PNU({assembled_pnu})와 좌표가 모순됩니다. 좌표가 있는 레코드는 "
                    "대표점으로 확정하지 않고 검토대상으로 분리합니다."
                ),
            )

        # ⑥ 실패 사유 구분 ---------------------------------------------------
        return self._classify_failure(assembled, assembled_pnu)

    def _classify_failure(
        self, assembled: AssembledPnu | None, assembled_pnu: str
    ) -> PnuResolution:
        if assembled is None:
            return PnuResolution(
                status="failed",
                reason=FAIL_NO_JIBUN,
                note=FAILURE_LABELS[FAIL_NO_JIBUN],
            )
        # 조립은 됐으나 지적도에 실재하지 않는다. 같은 본번의 다른 부번이 있는지로
        # 「지번 변경·합병·분할」과 「본번 자체 없음(폐쇄지번·오류)」을 가른다.
        prefix_exists = False
        if self.cadastral is not None:
            prefix_exists = self.cadastral.pnu_prefix_exists(assembled.bonbun_prefix)
        if prefix_exists:
            return PnuResolution(
                status="failed",
                reason=FAIL_JIBUN_NOT_IN_CADASTRAL,
                assembled_pnu=assembled_pnu,
                note=FAILURE_LABELS[FAIL_JIBUN_NOT_IN_CADASTRAL],
            )
        return PnuResolution(
            status="failed",
            reason=FAIL_BONBUN_NOT_IN_CADASTRAL,
            assembled_pnu=assembled_pnu,
            note=FAILURE_LABELS[FAIL_BONBUN_NOT_IN_CADASTRAL],
        )

    def _rep_of(self, parcel) -> Coordinates | None:
        if self.cadastral is None or parcel is None:
            return None
        return self.cadastral.representative_point(parcel)


# ── 배치 요약(보고·화면 note 용) ────────────────────────────────────────────


@dataclass
class ResolveSummary:
    """여러 원천 레코드 PNU 확정 결과 집계. 사유별로 나눠 보고할 수 있게 한다."""

    total: int = 0
    confirmed: int = 0
    review: int = 0
    failed: int = 0
    by_method: dict[str, int] = field(default_factory=dict)
    by_reason: dict[str, int] = field(default_factory=dict)
    confirmed_pnus: set[str] = field(default_factory=set)

    def add(self, resolution: PnuResolution) -> None:
        self.total += 1
        if resolution.status == "confirmed":
            self.confirmed += 1
            self.by_method[resolution.method] = (
                self.by_method.get(resolution.method, 0) + 1
            )
            if resolution.pnu:
                self.confirmed_pnus.add(resolution.pnu)
        elif resolution.status == "review":
            self.review += 1
            self.by_reason[resolution.reason] = (
                self.by_reason.get(resolution.reason, 0) + 1
            )
        else:
            self.failed += 1
            self.by_reason[resolution.reason] = (
                self.by_reason.get(resolution.reason, 0) + 1
            )

    @property
    def confirm_rate(self) -> float:
        return self.confirmed / self.total if self.total else 0.0


def resolve_records(
    resolver: PnuResolver,
    records: Iterable,
    *,
    address_attr: str = "address",
    pnu_attr: str = "pnu",
    coord_attr: str = "coordinates",
) -> tuple[ResolveSummary, list[tuple[object, PnuResolution]]]:
    """중립 레코드(LocalSourceRecord 등)를 순회하며 PNU 를 확정한다.

    반환: (집계, [(레코드, 확정결과)]). 레코드는 삭제하지 않는다 — 확정·검토·실패를
    모두 반환해 상위(배선·감사)가 격리·보고할 수 있게 한다.
    """

    summary = ResolveSummary()
    results: list[tuple[object, PnuResolution]] = []
    for record in records:
        address = getattr(record, address_attr, "") or ""
        pnu = getattr(record, pnu_attr, "") or ""
        coordinates = getattr(record, coord_attr, None)
        resolution = resolver.resolve(
            address=address, pnu=pnu, coordinates=coordinates
        )
        summary.add(resolution)
        results.append((record, resolution))
    return summary, results


# ── 조립 헬퍼(라우터 배선용) ─────────────────────────────────────────────────


def build_pnu_resolver(
    cadastral: CadastralBackend | None = None,
    geocoder: Geocoder | None = None,
    legal_dong: LegalDongIndex | None = None,
    address_pnu: AddressPnuAssembler | None = None,
) -> PnuResolver:
    """환경변수(LH_LEGAL_DONG_PATH) + 주입된 지적도로 해석기를 만든다.

    조립 경로는 법정동코드표(legal_dong)와 API(address_pnu) 중 하나면 된다. 둘 다
    없으면 조립이 비활성화되지만 입력 PNU·좌표 공간조인 경로는 살아 있다. 지적도
    백엔드까지 없으면 어떤 확정도 하지 않는다(빈 결과로 삼키지 않고 실패로 전파).
    다른 PC·미배포 환경에서 예외로 죽지 않는다.
    """

    if legal_dong is None:
        legal_dong = LegalDongIndex.from_env()
    if cadastral is None:
        from app.services.cadastral_local import CadastralLocalStore

        cadastral = CadastralLocalStore()
    return PnuResolver(
        legal_dong=legal_dong,
        cadastral=cadastral,
        geocoder=geocoder,
        address_pnu=address_pnu,
    )


__all__ = [
    "ENV_LEGAL_DONG",
    "AssembledPnu",
    "CadastralBackend",
    "FAILURE_ACTIONS",
    "FAILURE_LABELS",
    "FAIL_BONBUN_NOT_IN_CADASTRAL",
    "FAIL_JIBUN_NOT_IN_CADASTRAL",
    "FAIL_NO_JIBUN",
    "AddressPnuAssembler",
    "LegalDongIndex",
    "LegalDongMatch",
    "METHOD_ADDRESS_ASSEMBLY",
    "METHOD_CADASTRAL_REPRESENTATIVE_POINT",
    "METHOD_COORDINATE_SPATIAL_JOIN",
    "METHOD_INPUT_PNU",
    "PnuResolution",
    "PnuResolver",
    "REVIEW_ASSEMBLY_SPATIAL_MISMATCH",
    "REVIEW_COORDINATE_OUTSIDE_PARCEL",
    "REVIEW_INPUT_PNU_MISMATCH",
    "REVIEW_MULTIPLE_PARCELS",
    "RESOLVER_ALGORITHM_VERSION",
    "ResolveSummary",
    "assemble_pnu",
    "build_pnu_resolver",
    "normalize_address",
    "resolve_records",
    "validate_pnu",
]
