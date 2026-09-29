"""심사 결과 Excel(.xlsx) — LH 내부망 앱의 결과 내려받기와 같은 4시트 묶음.

시트: 종합요약 · 1차 상세 · 2차 상세 · 데이터 스냅샷.
한 건(심사표)과 일괄 심사가 같은 함수를 쓴다. 담당자 판단·메모는 브라우저에만
있으므로 프런트가 요청에 실어 보내고, 여기서는 1차 상세 시트에 그대로 적는다.
"""

from __future__ import annotations

import io
from datetime import UTC, datetime
from typing import Literal

from pydantic import BaseModel, Field

from app.hazard_review.rulebook import APPLICATION_TYPE_LABELS, HOUSING_TYPE_LABELS
from app.rules_config import OPTION_DEFS, is_relaxed, option_enabled
from app.screening.models import (
    ScreeningCriterion,
    ScreeningFacilityHit,
    ScreeningResult,
)


XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# 담당자 판단 5단계. 프런트 ScreeningSheet 의 선택지와 같은 키다.
JudgementState = Literal["unchecked", "confirmed", "evidence_requested", "false_positive", "excluded"]

JUDGEMENT_LABELS: dict[str, str] = {
    "unchecked": "미확인",
    "confirmed": "확인 완료",
    "evidence_requested": "추가 증빙 요청",
    "false_positive": "오탐",
    "excluded": "적용 제외",
}

MEASUREMENT_TIER_LABELS: dict[str, str] = {
    "front_door_parcel": "정문 필지",
    "front_door_point": "정문 좌표",
    "site_boundary": "시설 경계",
    "coordinate": "시설 좌표",
}


class ItemJudgement(BaseModel):
    """1차 항목 하나에 대한 담당자 판단과 메모."""

    state: JudgementState = "unchecked"
    memo: str = ""


class ExportSite(BaseModel):
    """Excel 한 줄(사업지 하나). 실패한 건은 result 없이 오류만 싣는다."""

    receipt_no: str = ""
    address: str = ""
    housing_type_label: str = ""
    application_type_label: str = ""
    result: ScreeningResult | None = None
    error: str = ""
    # 1차 항목 키(ScreeningExclusionItem.key) → 담당자 판단.
    judgements: dict[str, ItemJudgement] = Field(default_factory=dict)


class ScreeningExportRequest(BaseModel):
    """한 건(또는 여러 건) 심사 결과를 그대로 보내 Excel 을 받는다."""

    sites: list[ExportSite] = Field(min_length=1, max_length=500)


class BatchExportRequest(BaseModel):
    """일괄 심사 Excel. 결과는 서버에 있으므로 담당자 판단만 보낸다.

    screening_id → (1차 항목 키 → 판단).
    """

    judgements: dict[str, dict[str, ItemJudgement]] = Field(default_factory=dict)


def _score_text(criterion: ScreeningCriterion | None) -> object:
    if criterion is None:
        return ""
    if criterion.awarded is not None:
        return criterion.awarded
    return f"{criterion.awarded_min}~{criterion.awarded_max}"


def _meters(value: float | None) -> object:
    return "" if value is None else round(value)


def _stamp(value: datetime | None) -> str:
    if value is None:
        return ""
    return value.astimezone().strftime("%Y-%m-%d %H:%M")


def _measurement(hit: ScreeningFacilityHit) -> str:
    label = MEASUREMENT_TIER_LABELS.get(hit.measurement_tier, "")
    if label and hit.measurement_method:
        return f"{label} · {hit.measurement_method}"
    return label or hit.measurement_method or "시설 좌표"


def _site_labels(site: ExportSite) -> tuple[str, str, str, str]:
    """(접수번호, 소재지, 분류, 신청유형)."""

    result = site.result
    if result is None:
        return site.receipt_no, site.address, site.housing_type_label, site.application_type_label
    return (
        site.receipt_no or result.site.name,
        site.address or result.site.address,
        result.housing_type_label or HOUSING_TYPE_LABELS.get(result.housing_type, ""),
        result.application_type_label or APPLICATION_TYPE_LABELS.get(result.application_type, ""),
    )


def build_rows(sites: list[ExportSite]) -> dict[str, list[list[object]]]:
    summary: list[list[object]] = [[
        "접수번호", "소재지", "분류", "신청유형", "1차 판정", "1차 요지",
        "2차 총점", "2차 만점", "확정", "대중교통 접근성", "주거여건", "교육여건", "역세권 가점",
        "담당자 확인", "심사 ID", "오류",
    ]]
    stage_one: list[list[object]] = [[
        "접수번호", "소재지", "대분류", "항목", "기준거리(m)", "판정",
        "가장 가까운 시설", "거리(m)", "근거", "원천", "담당자 판단", "메모",
    ]]
    stage_two: list[list[object]] = [[
        "접수번호", "소재지", "평가항목", "점수", "만점", "시설군", "수집 상태",
        "시설명", "거리(m)", "측정기준", "배점 반영", "비고",
    ]]
    for site in sites:
        receipt, address, housing, application = _site_labels(site)
        result = site.result
        if result is None:
            summary.append([
                receipt, address, housing, application, "", "", "", "", "", "", "", "", "", "", "",
                site.error or "심사 결과 없음",
            ])
            continue
        two = result.stage_two
        scores = {c.key: c for c in two.criteria}
        judged = sum(
            1 for item in result.stage_one.items
            if site.judgements.get(item.key, ItemJudgement()).state != "unchecked"
        )
        summary.append([
            receipt, address, housing, application,
            result.stage_one.verdict_label, result.verdict_summary,
            two.living_score if two.determined and two.living_score is not None
            else f"{two.living_score_min}~{two.living_score_max}",
            two.living_maximum,
            "확정" if two.determined else "미확정",
            _score_text(scores.get("transit")),
            _score_text(scores.get("living")),
            _score_text(scores.get("education")),
            _score_text(two.bonus),
            f"{judged}/{len(result.stage_one.items)}",
            result.screening_id,
            site.error,
        ])
        for item in result.stage_one.items:
            nearest = min(item.facilities, key=lambda f: f.distance_m, default=None)
            judgement = site.judgements.get(item.key, ItemJudgement())
            stage_one.append([
                receipt, address, item.rule_label, item.label,
                item.threshold_m if item.threshold_m is not None else "미적용",
                item.outcome_label,
                nearest.name if nearest else "",
                _meters(nearest.distance_m if nearest else item.nearest_distance_m),
                item.reason, item.source_label,
                JUDGEMENT_LABELS[judgement.state], judgement.memo,
            ])
        for criterion in [*two.criteria, *([two.bonus] if two.bonus else [])]:
            for group in criterion.groups:
                if not group.hits:
                    stage_two.append([
                        receipt, address, criterion.label, _score_text(criterion), criterion.maximum,
                        group.label, group.state_label, "(해당 시설 없음)", "", "", "", group.note,
                    ])
                    continue
                for hit in group.hits:
                    stage_two.append([
                        receipt, address, criterion.label, _score_text(criterion), criterion.maximum,
                        group.label, group.state_label, hit.name, _meters(hit.distance_m),
                        _measurement(hit), "반영" if hit.counted else "제외",
                        hit.count_note or hit.front_door_notice,
                    ])
    return {"종합요약": summary, "1차 상세": stage_one, "2차 상세": stage_two}


def snapshot_rows(sites: list[ExportSite], generated_at: datetime) -> list[list[object]]:
    """데이터 스냅샷 시트. 위는 항목·값, 아래는 원천 장애 경고 표."""

    results = [site.result for site in sites if site.result is not None]
    packs = sorted({f"{r.rule_pack_id or '미지정'} v{r.rule_pack_version or '—'}" for r in results})
    screened = [r.created_at for r in results]
    snapshots = [r.data_snapshot for r in results if r.data_snapshot is not None]
    synced_min = min((s.localdata_synced_at_min for s in snapshots if s.localdata_synced_at_min), default=None)
    synced_max = max((s.localdata_synced_at_max for s in snapshots if s.localdata_synced_at_max), default=None)
    # 판정 옵션은 심사 시각에 기록한 값을 쓴다. 옛 결과라 기록이 없으면 지금 값으로 적는다.
    if snapshots:
        first = snapshots[0]
        options = [f"{o.label}: {'켬' if o.enabled else '끔'}" for o in first.options]
        relaxed = first.relaxed_2027
        option_basis = "심사 시각 기준"
    else:
        options = [f"{o.label}: {'켬' if option_enabled(o.key) else '끔'}" for o in OPTION_DEFS]
        relaxed = is_relaxed()
        option_basis = "내보낸 시각 기준(심사 결과에 기록 없음)"

    rows: list[list[object]] = [
        ["항목", "값"],
        ["생성 시각", _stamp(generated_at)],
        ["사업지 수", f"{len(sites)}건(심사 완료 {len(results)}건)"],
        ["심사 시각", f"{_stamp(min(screened))} ~ {_stamp(max(screened))}" if screened else ""],
        ["규칙팩 버전", ", ".join(packs)],
        ["2027 완화 기준", "적용" if relaxed else "미적용"],
        [f"적용 판정 옵션({option_basis})", " / ".join(options)],
        [
            "인허가(LOCALDATA) 캐시 동기화",
            f"{_stamp(synced_min)} ~ {_stamp(synced_max)}" if synced_min else "기록 없음",
        ],
        ["그 밖의 원천", "공공 API 실시간 조회(심사 시각 기준)"],
        [],
        ["원천 장애 경고"],
        ["접수번호", "단계", "원천", "내용"],
    ]
    alert_count = 0
    for site in sites:
        if site.result is None:
            continue
        receipt = _site_labels(site)[0]
        for alert in site.result.source_alerts:
            alert_count += 1
            rows.append([receipt, "1차" if alert.stage == "stage_one" else "2차", alert.source, alert.message])
    if alert_count == 0:
        rows.append(["", "", "", "경고 없음"])
    return rows


COLUMN_WIDTHS: dict[str, list[int]] = {
    "종합요약": [10, 36, 14, 12, 12, 40, 10, 8, 8, 10, 10, 10, 10, 10, 38, 30],
    "1차 상세": [10, 36, 18, 22, 11, 11, 28, 9, 50, 22, 14, 36],
    "2차 상세": [10, 36, 16, 8, 6, 16, 12, 30, 9, 24, 9, 40],
    "데이터 스냅샷": [30, 60, 24, 60],
}


def build_workbook(sites: list[ExportSite], generated_at: datetime | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    generated_at = generated_at or datetime.now(UTC)
    sheets = build_rows(sites)
    sheets["데이터 스냅샷"] = snapshot_rows(sites, generated_at)

    book = Workbook()
    book.remove(book.active)
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DCE6F1")
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(["" if value is None else value for value in row])
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
        if title != "데이터 스냅샷":
            sheet.freeze_panes = "A2"
        else:
            # 경고 표 머리행도 굵게.
            for row in sheet.iter_rows():
                if row and row[0].value in {"원천 장애 경고", "접수번호"}:
                    for cell in row:
                        cell.font = header_font
        for index, width in enumerate(COLUMN_WIDTHS.get(title, []), start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def xlsx_filename_header(prefix: str, generated_at: datetime) -> str:
    """Content-Disposition. 한글 파일명은 RFC 5987 로 싣고 ASCII 대체 이름도 준다."""

    from urllib.parse import quote

    stamp = generated_at.astimezone().strftime("%Y%m%d_%H%M")
    ascii_name = f"LH_screening_{stamp}.xlsx"
    utf8_name = quote(f"{prefix}_{stamp}.xlsx")
    return f"attachment; filename={ascii_name}; filename*=UTF-8''{utf8_name}"
