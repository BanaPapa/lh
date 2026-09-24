"""일괄 심사 — 신청자 엑셀(또는 CSV) 한 장을 올려 여러 사업지를 차례로 심사한다.

LH 제출본 앱의 「신청자 엑셀 올리기 → 검토 실행 → 결과 내려받기」 흐름을 그대로
따른다. 한 건 심사 경로(geocode → 필지 확보 → screen)를 바꾸지 않고 그 위에서
줄 세워 돌리므로, 한 건씩 돌린 결과와 일괄 결과가 같다.

파일 규칙 — LH 「신축매입 심사지 리스트 양식」(09/22 회의 결정 5)을 기준으로 한다.
  순번 · 차수 · 접수번호 · 접수일자 · 매도자명 · 분류 · 신청유형 · 지역 · 상세주소 · 경도 · 위도
  접수번호(없으면 순번·연번·번호) → 식별자(없으면 행 번호)
  분류(주택/주거용 오피스텔) · 신청유형(청년/일반/신혼/고령자) → 행별 유형. 비면 일괄 기본값
  지역 → 상세주소 앞에 붙인다(상세주소에 이미 있으면 붙이지 않는다)
  상세주소(소재지·주소도 받는다) → 필수. 다필지는 쉼표로 적는다(「외 N필지」만 있으면 대표 필지만 판정)
  경도·위도 → 있으면 지오코딩을 건너뛴다. 없어도 된다(주소로 좌표·PNU 를 잡는다)
  차수 · 접수일자 · 매도자명 → 판정에 쓰지 않고 결과에 그대로 싣는다. 그 밖의 열은 무시한다.
  이전 제출본 앱 열 이름(주택유형·건물유형·유형·pnu)도 그대로 받는다.
"""

from __future__ import annotations

import asyncio
import csv
import io
import logging
import re
import zipfile
from datetime import UTC, datetime
from typing import Literal
from uuid import uuid4

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, Field

from app.config import get_settings
from app.hazard_review.models import HazardParcelResolveRequest, HazardSite
from app.hazard_review.parcels import ParcelResolver
from app.hazard_review.router import get_parcel_resolver
from app.hazard_review.rulebook import (
    APPLICATION_TYPE_LABELS,
    HOUSING_TYPE_LABELS,
    ApplicationType,
    HousingType,
)
from app.models import Coordinates
from app.screening.models import ScreeningRequest, ScreeningResult
from app.screening.service import ScreeningService
from app.services.kakao import KakaoClient

logger = logging.getLogger("uvicorn.error")

router = APIRouter(prefix="/api/screening/batch", tags=["screening-batch"])

MAX_ROWS = 500
# 원천 API(카카오·TAGO·VWorld) 쿼터를 지키려고 한 번에 두 건까지만 돌린다.
CONCURRENCY = 2

# 헤더 이름 → 열 역할. LH 제출본 앱(analyze/page.tsx)과 같은 이름을 받는다.
KEY_ID = ("접수번호", "순번", "연번", "번호", "no")
KEY_REGION = ("지역", "시군", "시군구", "시도")
KEY_ADDR = ("상세주소", "소재지", "주소", "도로명주소", "지번주소")
KEY_PNU = ("pnu", "필지고유번호", "고유번호")
KEY_LNG = ("경도", "lng", "lon", "longitude", "x")
KEY_LAT = ("위도", "lat", "latitude", "y")
KEY_HOUSING = ("분류", "주택유형", "적용 건물유형", "적용건물유형", "건물유형")
KEY_APPLICATION = ("신청유형", "유형")
# 판정에 쓰지 않고 결과에 그대로 싣는 열.
KEY_EXTRAS: dict[str, tuple[str, ...]] = {
    "차수": ("차수",),
    "접수일자": ("접수일자", "접수일"),
    "매도자명": ("매도자명", "매도자", "신청자명", "신청자"),
}
EXTRA_KEYS = tuple(KEY_EXTRAS)

# 「외 N필지」만 있고 지번 목록이 없는 주소. 대표 필지만 판정되므로 담당자에게 알린다.
_OESILJI_NO_LIST = re.compile(r"외\s*\d+\s*필지(?!\s*\()")
TEMPLATE_COLUMNS = ("순번", "차수", "접수번호", "접수일자", "매도자명", "분류", "신청유형", "지역", "상세주소", "경도", "위도")


# ---------------------------------------------------------------------------
# 모델
# ---------------------------------------------------------------------------
class BatchRowInput(BaseModel):
    id: str
    address: str
    lat: float | None = None
    lng: float | None = None
    pnu: str = ""
    # 행별 유형. 비우면 일괄 기본값을 쓴다.
    housing_type: HousingType | None = None
    application_type: ApplicationType | None = None
    # 파일에 적힌 분류·신청유형 원문(인식 못 한 값을 화면에 그대로 보이기 위해).
    housing_label: str = ""
    application_label: str = ""
    # 차수·접수일자·매도자명 — 결과에 그대로 싣는다.
    extras: dict[str, str] = Field(default_factory=dict)
    # 담당자가 확인해야 할 점(분류 혼합·인식 불가·외 N필지 등).
    warnings: list[str] = Field(default_factory=list)


class BatchParseResponse(BaseModel):
    file_name: str
    sheet: str = ""
    header_row: int
    columns: dict[str, str] = Field(default_factory=dict)
    rows: list[BatchRowInput] = Field(default_factory=list)
    skipped: int = 0
    note: str = ""


class BatchStartRequest(BaseModel):
    rows: list[BatchRowInput] = Field(min_length=1, max_length=MAX_ROWS)
    housing_type: HousingType
    application_type: ApplicationType
    rule_pack_id: str = ""


RowState = Literal["queued", "running", "completed", "failed", "cancelled"]


class BatchCriterionScore(BaseModel):
    key: str
    label: str
    awarded: int | None
    awarded_min: int
    awarded_max: int
    maximum: int


class BatchRowStatus(BaseModel):
    id: str
    address: str
    housing_type: HousingType
    application_type: ApplicationType
    extras: dict[str, str] = Field(default_factory=dict)
    status: RowState = "queued"
    message: str = ""
    error: str = ""
    screening_id: str | None = None
    # 아래는 완료 후 채운다.
    site_address: str = ""
    parcel_count: int = 0
    resolve_note: str = ""
    verdict: str = ""
    verdict_label: str = ""
    verdict_summary: str = ""
    living_score: int | None = None
    living_score_min: int | None = None
    living_score_max: int | None = None
    living_maximum: int | None = None
    determined: bool = False
    criteria: list[BatchCriterionScore] = Field(default_factory=list)
    bonus: BatchCriterionScore | None = None


class BatchStatus(BaseModel):
    batch_id: str
    status: Literal["queued", "running", "completed", "cancelled"]
    created_at: datetime
    housing_type: HousingType
    application_type: ApplicationType
    total: int
    done: int = 0
    rows: list[BatchRowStatus] = Field(default_factory=list)


class BatchStart(BaseModel):
    batch_id: str


# ---------------------------------------------------------------------------
# 파일 읽기
# ---------------------------------------------------------------------------
def _norm(value: object) -> str:
    return re.sub(r"\s+", "", str(value or "")).strip().lower()


def _find_column(header: list[str], keys: tuple[str, ...]) -> int | None:
    """완전일치를 먼저 보고, 없으면 포함(「LH 파일 주소」→ 주소)으로 잡는다.

    「주소」와 「상세주소」가 함께 있을 때 엉뚱한 칸을 잡지 않도록 완전일치가 앞선다.
    x·y·no 같은 한두 글자 키는 포함 검사에 쓰지 않는다(「xlsx」 따위에 걸린다).
    """

    normalized = [_norm(cell) for cell in header]
    for key in keys:
        for index, cell in enumerate(normalized):
            if cell == key.lower():
                return index
    for key in keys:
        if len(key) < 3 and key.isascii():
            continue
        for index, cell in enumerate(normalized):
            if cell and key.lower() in cell:
                return index
    return None


def _to_float(value: object) -> float | None:
    if value is None or value == "":
        return None
    try:
        number = float(str(value).strip())
    except ValueError:
        return None
    return number if number == number else None  # NaN 방지


_HOUSE_TOKENS = ("주택", "다가구", "다세대", "아파트", "연립", "단독", "도시형")


def housing_type_from_label(text: str) -> HousingType | None:
    """「주거용오피스텔」·「다가구」·「아파트」처럼 원문 그대로 적힌 값을 코드로 옮긴다.

    도시형생활주택은 주택(공동주택)이다(최기헌 부장 2026-09-07). 「도시형생활주택 +
    주거용오피스텔」처럼 둘이 섞이면 양쪽 다 해당하는 건이라 여기서 정하지 않고
    None 을 돌려 담당자가 고르게 한다(is_mixed_housing_label 로 구분).
    """

    value = _norm(text)
    if not value:
        return None
    for code, label in HOUSING_TYPE_LABELS.items():
        if value == _norm(label) or value == code:
            return code  # type: ignore[return-value]
    is_officetel = "오피스텔" in value
    is_house = any(token in value for token in _HOUSE_TOKENS)
    if is_officetel and is_house:
        return None
    if is_officetel:
        return "officetel"
    if is_house:
        return "house"
    return None


def is_mixed_housing_label(text: str) -> bool:
    value = _norm(text)
    return "오피스텔" in value and any(token in value for token in _HOUSE_TOKENS)


def application_type_from_label(text: str) -> ApplicationType | None:
    value = _norm(text)
    if not value:
        return None
    for code, label in APPLICATION_TYPE_LABELS.items():
        if value == _norm(label) or value == code:
            return code  # type: ignore[return-value]
    if "청년" in value or "기숙사" in value:
        return "youth"
    if "신혼" in value or "신생아" in value:
        return "newlywed"
    if "다자녀" in value:
        return "multi_child"
    if "고령" in value:
        return "senior"
    if "일반" in value:
        return "general"
    return None


def _rows_from_table(table: list[list[object]], file_name: str, sheet: str = "") -> BatchParseResponse:
    """헤더 행을 찾아 열 역할을 정하고, 소재지가 있는 행만 골라 낸다."""

    header_index = None
    for index, row in enumerate(table[:40]):
        if _find_column([str(c or "") for c in row], KEY_ADDR) is not None:
            header_index = index
            break
    if header_index is None:
        raise HTTPException(
            status_code=400,
            detail="소재지 열을 찾지 못했습니다. 머리글에 「소재지」·「상세주소」·「주소」 중 하나가 있어야 합니다.",
        )
    header = [str(c or "") for c in table[header_index]]
    col = {
        "id": _find_column(header, KEY_ID),
        "region": _find_column(header, KEY_REGION),
        "address": _find_column(header, KEY_ADDR),
        "pnu": _find_column(header, KEY_PNU),
        "lng": _find_column(header, KEY_LNG),
        "lat": _find_column(header, KEY_LAT),
        "housing": _find_column(header, KEY_HOUSING),
        "application": _find_column(header, KEY_APPLICATION),
    }
    extra_cols = {
        key: index
        for key, keys in KEY_EXTRAS.items()
        if (index := _find_column(header, keys)) is not None
    }

    def cell(row: list[object], key: str) -> str:
        index = col[key]
        if index is None or index >= len(row):
            return ""
        value = row[index]
        if value is None:
            return ""
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()

    sheet_housing = housing_type_from_label(sheet)
    sheet_application = application_type_from_label(sheet)
    rows: list[BatchRowInput] = []
    skipped = 0
    for offset, raw in enumerate(table[header_index + 1 :], start=1):
        row = list(raw)
        address = cell(row, "address")
        if not address:
            skipped += 1
            continue
        region = cell(row, "region")
        if region and not address.startswith(region.split()[0]):
            address = f"{region} {address}"
        housing_label = cell(row, "housing")
        application_label = cell(row, "application")
        housing_type = housing_type_from_label(housing_label)
        application_type = application_type_from_label(application_label)
        # 시트 이름이 「청년」·「주택」처럼 유형이면(진용성 통합본 양식) 빈 칸의 기본값이 된다.
        if not housing_label and sheet_housing:
            housing_type, housing_label = sheet_housing, f"시트 「{sheet}」"
        if not application_label and sheet_application:
            application_type, application_label = sheet_application, f"시트 「{sheet}」"
        warnings: list[str] = []
        if housing_label and housing_type is None:
            if is_mixed_housing_label(housing_label):
                warnings.append(f"분류 「{housing_label}」 — 주택·주거용 오피스텔 양쪽에 해당해 하나를 골라야 합니다.")
            else:
                warnings.append(f"분류 「{housing_label}」 을(를) 인식하지 못했습니다 — 확인 필요.")
        if application_label and application_type is None:
            warnings.append(f"신청유형 「{application_label}」 을(를) 인식하지 못했습니다 — 확인 필요.")
        if _OESILJI_NO_LIST.search(address):
            warnings.append("「외 N필지」만 있고 지번 목록이 없어 대표 필지만 판정합니다. 지번을 쉼표로 적으면 합쳐서 판정합니다.")
        extras: dict[str, str] = {}
        for key, index in extra_cols.items():
            if index >= len(row) or row[index] in (None, ""):
                continue
            value = row[index]
            extras[key] = str(int(value)) if isinstance(value, float) and value.is_integer() else str(value).strip()
        rows.append(
            BatchRowInput(
                id=cell(row, "id") or str(offset),
                address=address,
                lat=_to_float(cell(row, "lat")),
                lng=_to_float(cell(row, "lng")),
                pnu=cell(row, "pnu"),
                housing_type=housing_type,
                application_type=application_type,
                housing_label=housing_label,
                application_label=application_label,
                extras=extras,
                warnings=warnings,
            )
        )
    if not rows:
        raise HTTPException(status_code=400, detail="소재지가 적힌 행이 없습니다.")
    if len(rows) > MAX_ROWS:
        raise HTTPException(status_code=400, detail=f"한 번에 {MAX_ROWS}건까지 올릴 수 있습니다({len(rows)}건).")
    return BatchParseResponse(
        file_name=file_name,
        sheet=sheet,
        header_row=header_index + 1,
        columns={
            **{key: header[index] for key, index in col.items() if index is not None},
            **{key: header[index] for key, index in extra_cols.items()},
        },
        rows=rows,
        skipped=skipped,
    )


def parse_upload(file_name: str, data: bytes) -> BatchParseResponse:
    lowered = file_name.lower()
    if lowered.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"엑셀 파일을 열지 못했습니다: {exc}") from exc
        # 소재지 열이 있는 첫 시트를 쓴다.
        last_error: HTTPException | None = None
        for sheet in workbook.worksheets:
            table = [list(row) for row in sheet.iter_rows(values_only=True)]
            try:
                return _rows_from_table(table, file_name, sheet.title)
            except HTTPException as exc:
                last_error = exc
        raise last_error or HTTPException(status_code=400, detail="읽을 시트가 없습니다.")
    if lowered.endswith(".csv"):
        for encoding in ("utf-8-sig", "cp949"):
            try:
                text = data.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise HTTPException(status_code=400, detail="CSV 인코딩을 읽지 못했습니다(UTF-8 또는 CP949).")
        table = [list(row) for row in csv.reader(io.StringIO(text))]
        return _rows_from_table(table, file_name)
    raise HTTPException(status_code=400, detail=".xlsx 또는 .csv 파일만 올릴 수 있습니다.")


# ---------------------------------------------------------------------------
# 실행
# ---------------------------------------------------------------------------
batches: dict[str, BatchStatus] = {}
batch_results: dict[tuple[str, str], ScreeningResult] = {}
batch_cancel: dict[str, asyncio.Event] = {}
_batch_tasks: set[asyncio.Task[None]] = set()


def _score(criterion) -> BatchCriterionScore:  # noqa: ANN001 — ScreeningCriterion
    return BatchCriterionScore(
        key=criterion.key,
        label=criterion.label,
        awarded=criterion.awarded,
        awarded_min=criterion.awarded_min,
        awarded_max=criterion.awarded_max,
        maximum=criterion.maximum,
    )


def _fill_row(row: BatchRowStatus, result: ScreeningResult, parcel_count: int, note: str) -> None:
    two = result.stage_two
    row.screening_id = result.screening_id
    row.site_address = result.site.address
    row.parcel_count = parcel_count
    row.resolve_note = note
    row.verdict = result.verdict
    row.verdict_label = result.verdict_label
    row.verdict_summary = result.verdict_summary
    row.living_score = two.living_score
    row.living_score_min = two.living_score_min
    row.living_score_max = two.living_score_max
    row.living_maximum = two.living_maximum
    row.determined = two.determined
    row.criteria = [_score(c) for c in two.criteria]
    row.bonus = _score(two.bonus) if two.bonus else None


async def run_row(
    row: BatchRowStatus,
    source: BatchRowInput,
    *,
    screening: ScreeningService,
    resolver: ParcelResolver,
    kakao: KakaoClient,
    rule_pack_id: str,
) -> ScreeningResult:
    """한 건 심사 경로 그대로: 좌표 → 필지 → 심사."""

    if source.lat is not None and source.lng is not None:
        coordinates = Coordinates(lat=source.lat, lng=source.lng)
    else:
        candidates = await kakao.geocode(source.address)
        if not candidates:
            raise RuntimeError("주소를 찾지 못했습니다. 소재지를 확인해 주세요.")
        coordinates = candidates[0].coordinates
    resolved = await resolver.resolve(
        HazardParcelResolveRequest(name=row.id, address=source.address, coordinates=coordinates)
    )
    representative = next((p for p in resolved.parcels if len(p.geometry) >= 4), None)
    request = ScreeningRequest(
        site=HazardSite(
            name=row.id,
            address=(representative.address if representative and representative.address else source.address),
            coordinates=coordinates,
            housing_type=row.housing_type,
            application_type=row.application_type,
            parcels=resolved.parcels,
        ),
        rule_pack_id=rule_pack_id,
        requested_by="lh-screening-batch",
    )
    result = await screening.screen(request)
    _fill_row(row, result, len(resolved.parcels), resolved.note)
    return result


async def run_batch(
    batch: BatchStatus,
    sources: list[BatchRowInput],
    *,
    screening: ScreeningService,
    resolver: ParcelResolver,
    kakao: KakaoClient,
    rule_pack_id: str,
    cancel: asyncio.Event,
    concurrency: int = CONCURRENCY,
) -> None:
    batch.status = "running"
    semaphore = asyncio.Semaphore(max(1, concurrency))

    async def one(row: BatchRowStatus, source: BatchRowInput) -> None:
        async with semaphore:
            if cancel.is_set():
                row.status = "cancelled"
                row.message = "중단"
                batch.done += 1
                return
            row.status = "running"
            row.message = "심사 중"
            try:
                result = await run_row(
                    row, source, screening=screening, resolver=resolver, kakao=kakao, rule_pack_id=rule_pack_id
                )
                batch_results[(batch.batch_id, row.id)] = result
                row.status = "completed"
                row.message = result.verdict_label
            except Exception as exc:  # 한 건 실패가 나머지를 막지 않는다.
                logger.warning("일괄 심사 실패 [%s] %s: %s", row.id, source.address, exc)
                row.status = "failed"
                row.error = str(exc) or type(exc).__name__
                row.message = "실패"
            finally:
                batch.done += 1

    await asyncio.gather(*(one(row, source) for row, source in zip(batch.rows, sources)))
    batch.status = "cancelled" if cancel.is_set() else "completed"


# ---------------------------------------------------------------------------
# 내보내기 — LH 제출본 앱과 같은 묶음(종합요약 · 1차 상세 · 2차 상세)
# ---------------------------------------------------------------------------
def _csv(rows: list[list[object]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    for row in rows:
        writer.writerow(["" if v is None else v for v in row])
    # 엑셀이 한글을 바로 읽도록 BOM 을 붙인다.
    return ("﻿" + buffer.getvalue()).encode("utf-8")


def _points(value: int | None) -> object:
    return "" if value is None else value


def export_rows(batch: BatchStatus) -> dict[str, list[list[object]]]:
    summary: list[list[object]] = [[
        "접수번호", *EXTRA_KEYS, "소재지", "분류", "신청유형", "상태", "1차 판정", "1차 요지",
        "2차 총점", "2차 만점", "확정", "산정 범위",
        "대중교통 접근성", "주거여건", "교육여건", "역세권 가점", "필지수", "필지 비고", "오류",
    ]]
    stage_one: list[list[object]] = [[
        "접수번호", "소재지", "항목", "세부", "기준거리(m)", "결과", "최근접(m)", "근거", "원천",
    ]]
    stage_two: list[list[object]] = [[
        "접수번호", "소재지", "평가항목", "점수", "만점", "채택 등급", "요건", "충족", "근거",
    ]]
    for row in batch.rows:
        scores = {c.key: c for c in row.criteria}
        summary.append([
            row.id, *(row.extras.get(key, "") for key in EXTRA_KEYS), row.address, HOUSING_TYPE_LABELS[row.housing_type],
            APPLICATION_TYPE_LABELS[row.application_type], row.message, row.verdict_label,
            row.verdict_summary, _points(row.living_score), _points(row.living_maximum),
            "확정" if row.determined else ("미확정" if row.status == "completed" else ""),
            "" if row.determined or row.living_score_min is None else f"{row.living_score_min}~{row.living_score_max}",
            _points(scores["transit"].awarded) if "transit" in scores else "",
            _points(scores["living"].awarded) if "living" in scores else "",
            _points(scores["education"].awarded) if "education" in scores else "",
            _points(row.bonus.awarded) if row.bonus else "",
            row.parcel_count or "", row.resolve_note, row.error,
        ])
        result = batch_results.get((batch.batch_id, row.id))
        if result is None:
            continue
        for item in result.stage_one.items:
            stage_one.append([
                row.id, row.address, item.rule_label, item.label,
                item.threshold_m if item.threshold_m is not None else "",
                item.outcome_label,
                round(item.nearest_distance_m) if item.nearest_distance_m is not None else "",
                item.reason, item.source_label,
            ])
        two = result.stage_two
        for criterion in [*two.criteria, *([two.bonus] if two.bonus else [])]:
            selected = next((t for t in criterion.tiers if t.selected), None)
            if selected is None:
                stage_two.append([
                    row.id, row.address, criterion.label, "", criterion.maximum,
                    criterion.tier_condition, "", "", criterion.note,
                ])
                continue
            if not selected.requirements:
                stage_two.append([
                    row.id, row.address, criterion.label, selected.points, criterion.maximum,
                    selected.condition, "", "", "",
                ])
            for req in selected.requirements:
                stage_two.append([
                    row.id, row.address, criterion.label, selected.points, criterion.maximum,
                    selected.condition, req.text,
                    "충족" if req.met else ("확인 불가" if req.met is None else "미충족"),
                    req.evidence,
                ])
    return {"종합요약": summary, "1차상세": stage_one, "2차상세": stage_two}


def export_zip(batch: BatchStatus) -> bytes:
    stamp = batch.created_at.astimezone().strftime("%Y-%m-%d")
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, rows in export_rows(batch).items():
            archive.writestr(f"일괄심사_{name}_{stamp}.csv", _csv(rows))
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# API
# ---------------------------------------------------------------------------
def _get_screening_service() -> ScreeningService:
    from app.screening.router import get_screening_service

    return get_screening_service()


def build_template() -> bytes:
    """LH 심사지 리스트 양식(09/22) — 머리글 한 줄과 예시 한 줄."""

    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "심사지 리스트"
    sheet.append(list(TEMPLATE_COLUMNS))
    sheet.append([1, "1001", "001", "26.04.13.", "", "주거용 오피스텔", "청년", "전주시", "완산구 효자동2가 1243-3", "", ""])
    widths = {"A": 6, "B": 8, "C": 10, "D": 12, "E": 12, "F": 16, "G": 10, "H": 10, "I": 40, "J": 12, "K": 12}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@router.get("/template")
async def download_template() -> Response:
    # 한글 파일명은 RFC 5987 로 싣는다.
    filename = "UTF-8''LH_%EC%8B%AC%EC%82%AC%EC%A7%80_%EB%A6%AC%EC%8A%A4%ED%8A%B8_%EC%96%91%EC%8B%9D.xlsx"
    return Response(
        content=build_template(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*={filename}"},
    )


@router.post("/parse", response_model=BatchParseResponse)
async def parse_batch_file(file: UploadFile = File(...)) -> BatchParseResponse:
    data = await file.read()
    if len(data) > 20 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="파일이 20MB 를 넘습니다.")
    return parse_upload(file.filename or "upload", data)


@router.post("", response_model=BatchStart)
async def start_batch(
    payload: BatchStartRequest,
    screening: ScreeningService = Depends(_get_screening_service),
    resolver: ParcelResolver = Depends(get_parcel_resolver),
) -> BatchStart:
    from app.screening.router import _await_warmup

    await _await_warmup(screening)
    batch_id = str(uuid4())
    batch = BatchStatus(
        batch_id=batch_id,
        status="queued",
        created_at=datetime.now(UTC),
        housing_type=payload.housing_type,
        application_type=payload.application_type,
        total=len(payload.rows),
        rows=[
            BatchRowStatus(
                id=source.id,
                address=source.address,
                housing_type=source.housing_type or payload.housing_type,
                application_type=source.application_type or payload.application_type,
                extras=source.extras,
            )
            for source in payload.rows
        ],
    )
    batches[batch_id] = batch
    cancel = asyncio.Event()
    batch_cancel[batch_id] = cancel
    kakao = KakaoClient(get_settings().kakao_rest_api_key)
    task = asyncio.create_task(
        run_batch(
            batch, payload.rows, screening=screening, resolver=resolver, kakao=kakao,
            rule_pack_id=payload.rule_pack_id, cancel=cancel,
        )
    )
    _batch_tasks.add(task)
    task.add_done_callback(_batch_tasks.discard)
    return BatchStart(batch_id=batch_id)


def _batch_or_404(batch_id: str) -> BatchStatus:
    batch = batches.get(batch_id)
    if batch is None:
        raise HTTPException(status_code=404, detail="일괄 심사 작업을 찾을 수 없습니다.")
    return batch


@router.get("/{batch_id}", response_model=BatchStatus)
async def get_batch(batch_id: str) -> BatchStatus:
    return _batch_or_404(batch_id)


@router.post("/{batch_id}/cancel", response_model=BatchStatus)
async def cancel_batch(batch_id: str) -> BatchStatus:
    batch = _batch_or_404(batch_id)
    event = batch_cancel.get(batch_id)
    if event:
        event.set()
    if batch.status in {"queued", "running"}:
        batch.status = "cancelled"
    return batch


@router.get("/{batch_id}/rows/{row_id}/result", response_model=ScreeningResult)
async def get_batch_row_result(batch_id: str, row_id: str) -> ScreeningResult:
    _batch_or_404(batch_id)
    result = batch_results.get((batch_id, row_id))
    if result is None:
        raise HTTPException(status_code=404, detail="이 건의 심사 결과가 아직 없습니다.")
    return result


@router.get("/{batch_id}/export")
async def export_batch(batch_id: str) -> Response:
    batch = _batch_or_404(batch_id)
    stamp = batch.created_at.astimezone().strftime("%Y-%m-%d")
    return Response(
        content=export_zip(batch),
        media_type="application/zip",
        headers={"Content-Disposition": f"attachment; filename=LH_batch_{stamp}.zip"},
    )
