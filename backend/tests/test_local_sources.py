"""로컬 원천 적재 모듈 단위 테스트.

실파일 없이 돈다. 픽스처는 작게 직접 만든다(인코딩 혼재·좌표 결측·상태 필터·
PNU 매칭·파일 부재를 모두 자체 픽스처로 재현). 실데이터 검증은 별도로 하고
여기서는 하지 않는다.
"""

from __future__ import annotations

import csv
import io
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.services.local_sources import (
    LocalSourcesConfig,
    classify_status,
    coerce_lat_lng,
    factory_registry_pnus,
    load_all,
    load_cng_stations,
    load_factory_prohibit_list,
    load_factory_registry,
    load_gas_product_manufacturers,
    load_gas_stations,
    load_lpg_charging_stations,
    load_standard_source,
    match_by_pnu,
    pnus_of,
)


# ── 픽스처 빌더 ──────────────────────────────────────────────────────────────


def _write_csv(path: Path, header: list[str], rows: list[list[str]], encoding: str) -> None:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(header)
    writer.writerows(rows)
    path.write_bytes(buffer.getvalue().encode(encoding))


def _write_factory_registry_xlsx(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(
        ["순번", "회사명", "공장주소(지번)", "공장주소(도로명)", "업종번호", "업종명", "산업단지명"]
    )
    worksheet.append(
        ["1", "가나공장", "전북 군산시 소룡동 1-1", "전북 군산시 가도로 1", "20312, 20313", "복합비료 제조업", "군산산단"]
    )
    worksheet.append(
        ["2", "다라공장", "전북 김제시 용지면 2-2", "전북 김제시 용지로 2", "68112", "임대업", ""]
    )
    workbook.save(path)


def _write_standard_xlsx(path: Path) -> None:
    workbook = Workbook()
    hazard = workbook.active
    hazard.title = "유해시설"
    header = [
        "facility_id", "dataset_id", "major_category", "middle_category",
        "facility_name", "address", "city_name", "location_basis", "pnu",
        "longitude", "latitude", "geometry_ref", "pnu_source", "pnu_method",
        "coordinate_source", "coordinate_method", "cadastral_match_status",
        "location_status",
    ]
    hazard.append(header)
    # dataset_id=2 middle=공장 (PNU 보유) — factoryON PNU 정본
    hazard.append([
        "10", "2", "공장 또는 위험물 저장 및 처리 시설", "공장",
        "표준공장A", "전북 군산시 소룡동 1-1", "군산시", "PNU", "PNU-AAA",
        "126.6", "35.9", None, "s", "m", "s", "m", "MATCHED", "SUCCESS",
    ])
    # dataset_id=2 middle=위험물 — 공장 필터에는 안 잡혀야 한다
    hazard.append([
        "11", "2", "공장 또는 위험물 저장 및 처리 시설", "위험물 저장 및 처리 시설",
        "위험물B", "전북 군산시 오식도동 5-1", "군산시", "PNU", "PNU-BBB",
        "126.55", "35.96", None, "s", "m", "s", "m", "MATCHED", "SUCCESS",
    ])
    amenity = workbook.create_sheet("생활편의시설")
    amenity.append(header)
    amenity.append([
        "500", "16", "상업시설", "전통시장",
        "고창시장", "전북 고창군 고창읍 1", "고창군", "PNU", "PNU-MKT",
        "126.7", "35.43", None, "s", "m", "s", "m", "MATCHED", "SUCCESS",
    ])
    workbook.save(path)


@pytest.fixture
def raw_dir(tmp_path: Path) -> Path:
    """유해시설/·편의시설/ 하위를 갖춘 원본 디렉터리 픽스처."""

    hazard = tmp_path / "유해시설"
    hazard.mkdir()

    # 고시업종 목록 (utf-8-sig)
    _write_csv(
        hazard / "factory_prohibit_list.csv",
        ["일련번호", "한국표준산업분류번호", "업종"],
        [["1", "20312", "복합비료 제조업"], ["2", "10110", "도축업"]],
        "utf-8-sig",
    )
    # 주유소: 변동사유 이력, 폐업/휴업/신규 섞임 (utf-8-sig)
    _write_csv(
        hazard / "11_gas_stations.csv",
        ["연도", "변동사유발생연월일", "판매업종류", "구분", "업체명", "소재지"],
        [
            ["2015", "2015.1.7", "주유소", "휴업", "휴업주유소", "전북 부안군 보안면 1"],
            ["2015", "2015.1.28", "주유소", "신규등록", "신규주유소", "전북 전주시 덕진구 2"],
            ["2016", "2016.3.3", "주유소", "폐업", "폐업주유소", "전북 군산시 3"],
            ["2016", "2016.4.4", "주유소", "등록취소", "취소주유소", "전북 익산시 4"],
        ],
        "utf-8-sig",
    )
    # LPG 충전소: LAT/LOT 컬럼명 혼동 + 좌표 결측 1행 (utf-8-sig)
    _write_csv(
        hazard / "13_lpg_charging_stations.csv",
        ["SECT_NM", "LOT", "MGT_NM", "TELNO", "BSES_NM", "LAT", "ADDR"],
        [
            ["전북 고창군", "126.69", "자동차", "063", "광진충전소", "35.43", "전북 고창군 중앙로 152"],
            ["전북 완주군", "", "자동차", "063", "좌표없음충전소", "", "전북 완주군 5"],
        ],
        "utf-8-sig",
    )
    # CNG: cp949, 위도/경도 보유
    _write_csv(
        hazard / "25_cng_stations.csv",
        ["순번", "행정구역", "지사", "시설명", "우편", "주소", "위도", "경도"],
        [["28", "전북 전주시", "전북본부", "팔복CNG", "54845", "전북 전주시 1", "35.84", "127.06"]],
        "cp949",
    )
    _write_factory_registry_xlsx(hazard / "03_06_09_factory_registry.xlsx")

    amenity = tmp_path / "편의시설"
    amenity.mkdir()
    return tmp_path


@pytest.fixture
def standard_path(tmp_path: Path) -> Path:
    path = tmp_path / "facilities.xlsx"
    _write_standard_xlsx(path)
    return path


# ── classify_status ─────────────────────────────────────────────────────────


def test_classify_excluded_keywords():
    assert classify_status("폐업", has_status_field=True)[0] == "excluded"
    assert classify_status("등록취소", has_status_field=True)[0] == "excluded"
    assert classify_status("말소", has_status_field=True)[0] == "excluded"


def test_classify_review_keeps_business_under_review():
    status_class, reason = classify_status("휴업", has_status_field=True)
    assert status_class == "review"
    assert "휴업" in reason


def test_classify_blank_status_is_review_when_field_exists():
    status_class, reason = classify_status("", has_status_field=True)
    assert status_class == "review"
    assert "공란" in reason


def test_classify_blank_is_active_when_no_status_field():
    # 상태 컬럼 자체가 없는 대장은 공란을 검토로 보지 않는다.
    assert classify_status("", has_status_field=False)[0] == "active"


# ── coerce_lat_lng ──────────────────────────────────────────────────────────


def test_coerce_lat_lng_resolves_by_value_range():
    # 이름과 무관하게 값으로 판별: 35.x=lat, 126.x=lng
    assert coerce_lat_lng("35.43", "126.69") == (35.43, 126.69)
    # 순서를 바꿔 넣어도 동일하게 판별한다.
    assert coerce_lat_lng("126.69", "35.43") == (35.43, 126.69)


def test_coerce_lat_lng_missing_returns_none():
    assert coerce_lat_lng("", "") == (None, None)


# ── 인코딩 혼재 ─────────────────────────────────────────────────────────────


def test_cp949_source_decodes(raw_dir: Path):
    config = LocalSourcesConfig(raw_path=raw_dir)
    result = load_cng_stations(config)
    assert result.loaded
    assert result.records[0].name == "팔복CNG"
    assert result.records[0].coordinates is not None


def test_utf8_sig_source_decodes(raw_dir: Path):
    config = LocalSourcesConfig(raw_path=raw_dir)
    result = load_lpg_charging_stations(config)
    assert result.loaded
    names = {r.name for r in result.records}
    assert "광진충전소" in names


# ── 영업상태 필터 ────────────────────────────────────────────────────────────


def test_gas_station_status_filter(raw_dir: Path):
    config = LocalSourcesConfig(raw_path=raw_dir)
    result = load_gas_stations(config)
    kept_names = {r.name for r in result.records}
    # 폐업·등록취소는 판정 투입 전 제외
    assert "폐업주유소" not in kept_names
    assert "취소주유소" not in kept_names
    assert result.excluded_count == 2
    # 휴업은 제외하지 않고 검토로 유지
    assert "휴업주유소" in kept_names
    review = {r.name for r in result.records if r.status_class == "review"}
    assert "휴업주유소" in review
    assert "신규주유소" in kept_names


# ── 좌표 결측 격리 ──────────────────────────────────────────────────────────


def test_missing_coordinates_are_quarantined_not_deleted(raw_dir: Path):
    config = LocalSourcesConfig(raw_path=raw_dir)
    result = load_lpg_charging_stations(config)
    quarantined = {r.name for r in result.quarantined}
    assert "좌표없음충전소" in quarantined
    # 삭제하지 않고 사유를 남긴다.
    q = next(r for r in result.quarantined if r.name == "좌표없음충전소")
    assert q.coordinates is None
    assert q.quarantine_reason
    # 좌표 있는 행은 판정 후보로 유지된다.
    assert any(r.name == "광진충전소" for r in result.records)


# ── factoryON / 고시업종 ────────────────────────────────────────────────────


def test_factory_registry_parses_ksic(raw_dir: Path):
    config = LocalSourcesConfig(raw_path=raw_dir)
    result = load_factory_registry(config)
    assert result.loaded
    first = next(r for r in result.records if r.name == "가나공장")
    assert first.ksic_codes == ("20312", "20313")
    # 원본에는 좌표·PNU 가 없다.
    assert first.coordinates is None
    assert first.pnu == ""


def test_factory_prohibit_list_codes(raw_dir: Path):
    config = LocalSourcesConfig(raw_path=raw_dir)
    prohibit = load_factory_prohibit_list(config)
    assert prohibit.loaded
    assert "20312" in prohibit.ksic_codes
    assert "10110" in prohibit.ksic_codes


# ── PNU 매칭키 ──────────────────────────────────────────────────────────────


def test_factory_registry_pnus_from_standard(standard_path: Path):
    config = LocalSourcesConfig(standard_path=standard_path)
    result = factory_registry_pnus(config)
    assert result.loaded
    # dataset_id=2 middle=공장 만 잡히고 위험물은 빠진다.
    assert result.pnus == frozenset({"PNU-AAA"})
    assert "PNU-BBB" not in result.pnus


def test_factory_registry_pnus_blank_pnu_reports_not_loaded(tmp_path: Path):
    # 파일은 있으나 공장 행의 PNU 가 전부 공란이면, 파일 존재를 loaded 로 보면
    # 좌표·PNU 해석 가능한 공장 스냅샷이 없는데도 AND 대조를 '가능'으로 오인한다.
    # 유효 PNU 0개면 loaded=False 여야 한다.
    path = tmp_path / "no_pnu_factory.xlsx"
    workbook = Workbook()
    hazard = workbook.active
    hazard.title = "유해시설"
    header = [
        "facility_id", "dataset_id", "major_category", "middle_category",
        "facility_name", "address", "city_name", "location_basis", "pnu",
        "longitude", "latitude", "geometry_ref", "pnu_source", "pnu_method",
        "coordinate_source", "coordinate_method", "cadastral_match_status",
        "location_status",
    ]
    hazard.append(header)
    # dataset_id=2 middle=공장 이지만 pnu 컬럼이 공란.
    hazard.append([
        "10", "2", "공장 또는 위험물 저장 및 처리 시설", "공장",
        "PNU없는공장", "전북 군산시 소룡동 1-1", "군산시", "PNU", "",
        "126.6", "35.9", None, "s", "m", "s", "m", "UNMATCHED", "SUCCESS",
    ])
    workbook.save(path)

    config = LocalSourcesConfig(standard_path=path)
    result = factory_registry_pnus(config)
    assert result.loaded is False
    assert result.pnus == frozenset()
    assert "유효 PNU" in (result.note or "")


def test_standard_factory_filter_excludes_hazardous_material(standard_path: Path):
    config = LocalSourcesConfig(standard_path=standard_path)
    factory = load_standard_source(config, "factory_standard")
    assert {r.name for r in factory.records} == {"표준공장A"}
    hazmat = load_standard_source(config, "hazardous_material_standard")
    assert {r.name for r in hazmat.records} == {"위험물B"}


def test_match_by_pnu_uses_exact_key(standard_path: Path):
    config = LocalSourcesConfig(standard_path=standard_path)
    factory = load_standard_source(config, "factory_standard")
    hazmat = load_standard_source(config, "hazardous_material_standard")
    # 서로 다른 PNU 라 매칭 없음(문자열 유사매칭으로 억지로 붙지 않는다).
    assert match_by_pnu(factory.records, hazmat.records) == []
    # 같은 PNU 를 넣으면 매칭된다.
    assert pnus_of(factory.records) == frozenset({"PNU-AAA"})


def test_traditional_markets_from_standard(standard_path: Path):
    config = LocalSourcesConfig(standard_path=standard_path)
    result = load_standard_source(config, "traditional_markets")
    assert result.loaded
    assert result.records[0].name == "고창시장"
    assert result.records[0].pnu == "PNU-MKT"
    assert result.records[0].coordinates is not None


# ── 파일 부재 → 미적재 ──────────────────────────────────────────────────────


def test_missing_raw_path_reports_not_loaded():
    config = LocalSourcesConfig(raw_path=None)
    result = load_factory_registry(config)
    assert result.loaded is False
    assert result.records == ()
    assert "미적재" in result.note


def test_missing_standard_path_reports_not_loaded():
    config = LocalSourcesConfig(standard_path=Path("does-not-exist.xlsx"))
    result = factory_registry_pnus(config)
    assert result.loaded is False


def test_missing_file_in_existing_dir(raw_dir: Path):
    # 디렉터리는 있지만 gas_product 파일은 픽스처에 없다 → 미적재.
    from app.services.local_sources import load_gas_product_manufacturers

    config = LocalSourcesConfig(raw_path=raw_dir)
    result = load_gas_product_manufacturers(config)
    assert result.loaded is False


# ── zip 원천 지원 ───────────────────────────────────────────────────────────


def test_reads_from_zip(raw_dir: Path, tmp_path: Path):
    zip_path = tmp_path / "sources.zip"
    with zipfile.ZipFile(zip_path, "w") as archive:
        for path in raw_dir.rglob("*"):
            if path.is_file():
                archive.write(path, path.relative_to(raw_dir).as_posix())
    config = LocalSourcesConfig(raw_path=zip_path)
    result = load_cng_stations(config)
    assert result.loaded
    assert result.records[0].name == "팔복CNG"


# ── load_all 통합 ───────────────────────────────────────────────────────────


def test_load_all_marks_missing_and_present(raw_dir: Path, standard_path: Path):
    config = LocalSourcesConfig(raw_path=raw_dir, standard_path=standard_path)
    results = load_all(config)
    assert results["factory_registry"].loaded is True
    assert results["cng_stations"].loaded is True
    assert results["traditional_markets"].loaded is True
    # 픽스처에 없는 원천은 미적재로 그대로 포함된다.
    assert results["gas_product_manufacturers"].loaded is False


def test_to_localdata_record_bridge(raw_dir: Path):
    config = LocalSourcesConfig(raw_path=raw_dir)
    result = load_cng_stations(config)
    bridged = result.records[0].to_localdata_record()
    assert bridged is not None
    assert bridged.dataset_key == "cng_stations"
    assert bridged.name == "팔복CNG"


# ── 좌표 정비 (docs/hazards H-03 §6-2 · H-02-나 §6-4 · H-02-아 §6-3) ────────


def _cfg_with(tmp_path: Path, filename: str, header: list[str], rows: list[list[str]],
              encoding: str = "utf-8-sig") -> LocalSourcesConfig:
    hazard = tmp_path / "유해시설"
    hazard.mkdir(exist_ok=True)
    _write_csv(hazard / filename, header, rows, encoding)
    return LocalSourcesConfig(raw_path=tmp_path)


CNG_HEADER = ["순번", "행정구역", "지사", "시설명", "우편", "주소", "위도", "경도"]


def test_cng_verified_coordinates_replace_original(tmp_path: Path):
    """H-03 §6-2: 원본 좌표 8건 전부 714~3,779m 오류 → VWorld·카카오 실측값으로 교체."""

    config = _cfg_with(
        tmp_path, "25_cng_stations.csv", CNG_HEADER,
        [["28", "전북 전주시 덕진구", "전북본부", "(유)호남고속팔복CNG충전소", "54845",
          "전북 전주시 덕진구 신복천변로 32", "35.845892", "127.065276"]],
        "cp949",
    )
    result = load_cng_stations(config)
    record = result.records[0]
    assert record.coordinates is not None
    assert abs(record.coordinates.lat - 35.857275) < 1e-6
    assert abs(record.coordinates.lng - 127.104786) < 1e-6
    assert not record.quarantined
    # 원본 값과 교체 근거를 남긴다(감사 추적).
    assert record.extra["original_lat"] == "35.845892"
    assert record.extra["original_lng"] == "127.065276"
    assert "H-03" in record.extra["coordinate_basis"]


def test_cng_verified_coordinates_match_by_name_ignoring_spaces(tmp_path: Path):
    config = _cfg_with(
        tmp_path, "25_cng_stations.csv", CNG_HEADER,
        [["53", "완주군", "전북본부", "현대자동차(주) 전주공장 CNG제1충전소", "",
          "완주군 봉동읍 완주산단5로 163", "35.952529", "127.126486"]],
        "cp949",
    )
    record = load_cng_stations(config).records[0]
    assert abs(record.coordinates.lat - 35.945800) < 1e-6
    assert abs(record.coordinates.lng - 127.138686) < 1e-6


def test_cng_unknown_station_keeps_original_coordinates(tmp_path: Path):
    config = _cfg_with(
        tmp_path, "25_cng_stations.csv", CNG_HEADER,
        [["99", "전북 전주시", "전북본부", "신규CNG충전소", "", "전북 전주시 1",
          "35.84", "127.06"]],
        "cp949",
    )
    record = load_cng_stations(config).records[0]
    assert abs(record.coordinates.lat - 35.84) < 1e-9
    assert "coordinate_basis" not in record.extra


LPG_HEADER = ["SECT_NM", "LOT", "MGT_NM", "TELNO", "BSES_NM", "LAT", "ADDR"]


def test_seoul_city_hall_default_coordinate_is_quarantined(tmp_path: Path):
    """H-02-나 §6-4 ①: 지오코딩 실패 기본값(서울시청 37.5665, 126.978)이 섞여 있다."""

    config = _cfg_with(
        tmp_path, "13_lpg_charging_stations.csv", LPG_HEADER,
        [["전북 완주군", "126.978", "자동차", "063", "현대자동차(주)전주공장", "37.5665",
          "전북 완주군 봉동읍 완주산단5로 163"]],
    )
    result = load_lpg_charging_stations(config)
    assert not result.records
    q = result.quarantined[0]
    assert q.coordinates is None
    assert "기본값" in q.quarantine_reason


def test_out_of_jeonbuk_coordinate_is_quarantined(tmp_path: Path):
    """H-02-나 §6-4 ②: 경도 128.7~128.9(경상도 권역) 좌표는 적재 단계에서 격리."""

    config = _cfg_with(
        tmp_path, "13_lpg_charging_stations.csv", LPG_HEADER,
        [["전북 김제시", "128.9018102", "자동차", "063", "우리LPG충전소", "35.9426894",
          "전북 김제시 금구면 1"]],
    )
    result = load_lpg_charging_stations(config)
    assert not result.records
    assert "전북 범위 밖" in result.quarantined[0].quarantine_reason


def test_duplicate_coordinates_across_sigungu_flagged_for_review(tmp_path: Path):
    """H-02-나 §6-4 ③: 동명 업소 좌표 복사 — 전주 건 좌표가 정읍 건에 그대로 들어 있다.

    어느 쪽이 맞는지 적재 단계에서는 알 수 없으므로 삭제·격리하지 않고 둘 다
    검토(review)로 낮춰 확정 판정에 쓰이지 않게 한다.
    """

    config = _cfg_with(
        tmp_path, "13_lpg_charging_stations.csv", LPG_HEADER,
        [
            ["전북 전주시", "127.130375", "자동차", "063", "호남가스충전소", "35.839469",
             "전북 전주시 덕진구 금암동 1"],
            ["전북 정읍시", "127.130375", "자동차", "063", "호남가스충전소", "35.839469",
             "전북 정읍시 정읍남로 1309"],
            ["전북 고창군", "126.69", "자동차", "063", "광진충전소", "35.43",
             "전북 고창군 중앙로 152"],
        ],
    )
    result = load_lpg_charging_stations(config)
    by_name = {}
    for record in result.records:
        by_name.setdefault(record.name, []).append(record)
    assert len(by_name["호남가스충전소"]) == 2
    assert all(r.status_class == "review" for r in by_name["호남가스충전소"])
    assert all("좌표 복사" in r.review_reason for r in by_name["호남가스충전소"])
    assert by_name["광진충전소"][0].status_class == "active"


def test_gas_product_manufacturers_dedupe_business_rows(tmp_path: Path):
    """H-02-아 §6-3: 28행 CSV 는 「업소 × 생산품목」 행이라 고유 업소로 접는다."""

    config = _cfg_with(
        tmp_path, "22_gas_product_manufacturers.csv",
        ["행정구역", "업소명", "소재지", "생산품목", "영업상태", "법구분"],
        [
            ["군산시", "(주)에쎈테크", "전북 군산시 산단로 1", "압력용기", "영업", "고법"],
            ["군산시", "(주)에쎈테크", "전북 군산시 산단로 1", "저장탱크", "영업", "고법"],
            ["군산시", "(주)에쎈테크", "전북 군산시 산단로 1", "배관용밸브", "영업", "고법"],
            ["완주군", "(유)금호주방", "전북 완주군 봉동읍 2", "업무용대형연소기", "영업", "액법"],
        ],
    )
    result = load_gas_product_manufacturers(config)
    assert len(result.records) == 2
    essen = next(r for r in result.records if r.name == "(주)에쎈테크")
    assert essen.category == "압력용기·저장탱크·배관용밸브"
    assert essen.extra["merged_rows"] == "3"


def test_bundle_records_actual_source_files(raw_dir: Path, standard_path: Path):
    """번들이 로컬 후보의 실제 적재 출처(파일명·origin)를 보존하는지.

    등록공장 후보는 표준본(facilities.xlsx)에서, CNG 후보는 원본
    (25_cng_stations.csv)에서 온다. 칩 detail 이 이 값을 써야 표시가 안 틀린다.
    """

    from app.services.local_wiring import LocalSourceFile, build_local_sources_bundle

    config = LocalSourcesConfig(standard_path=standard_path, raw_path=raw_dir)
    bundle = build_local_sources_bundle(config)

    factory = bundle.source_files["factory_registry"]
    assert isinstance(factory, LocalSourceFile)
    assert factory.filename == "facilities.xlsx"
    assert factory.origin == "standard"
    # 표준본은 파일 공유라 짧은 구분이 붙는다(원본 파일명 하드코딩 금지).
    assert factory.qualifier
    assert "facilities.xlsx" in factory.chip_detail()

    cng = bundle.source_files["cng_stations"]
    assert cng.filename == "25_cng_stations.csv"
    assert cng.origin == "raw"
    assert cng.chip_detail() == "25_cng_stations.csv"
