"""일괄 심사 — 파일 읽기(열 이름·지역 접두·행별 유형), 줄 세워 돌리기, 내보내기."""

from __future__ import annotations

import asyncio
import io
import os
from datetime import UTC, datetime

import pytest
from fastapi import HTTPException
from openpyxl import Workbook

os.environ["DEMO_MODE"] = "true"

from app.hazard_review.models import HazardParcel, HazardParcelResolveResponse
from app.models import Coordinates, GeocodeCandidate
from app.screening.batch import (
    BatchRowInput,
    BatchRowStatus,
    BatchStatus,
    application_type_from_label,
    export_rows,
    housing_type_from_label,
    parse_upload,
    run_batch,
)
from app.screening.models import ScreeningRequest

from tests.test_screening_scorebook import (  # noqa: E402
    FakeAmenityCollector,
    FakeHazardService,
    category,
    collections,
)
from app.screening.service import ScreeningService

CENTER = Coordinates(lat=35.8152, lng=127.104)


def _xlsx(rows: list[list[object]]) -> bytes:
    book = Workbook()
    sheet = book.active
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_parse_xlsx_finds_header_and_prefixes_region() -> None:
    data = _xlsx(
        [
            ["신축매입약정 신청 현황", None, None],
            ["접수번호", "지역", "상세주소", "위도", "경도", "유형", "건물유형"],
            ["002", "전주시", "완산구 효자동2가 1243-3", 35.8152, 127.104, "청년·기숙사형", "주거용오피스텔"],
            ["004", "전주시", "전주시 덕진구 송천동1가 626-74", None, None, "일반", "다가구"],
            [None, None, None, None, None, None, None],
        ]
    )
    parsed = parse_upload("신청자.xlsx", data)

    assert parsed.header_row == 2
    assert parsed.skipped == 1
    assert [r.id for r in parsed.rows] == ["002", "004"]
    # 지역 열은 소재지 앞에 붙이되, 이미 들어 있으면 두 번 붙이지 않는다.
    assert parsed.rows[0].address == "전주시 완산구 효자동2가 1243-3"
    assert parsed.rows[1].address == "전주시 덕진구 송천동1가 626-74"
    assert parsed.rows[0].lat == 35.8152 and parsed.rows[1].lat is None
    assert parsed.rows[0].housing_type == "officetel" and parsed.rows[0].application_type == "youth"
    assert parsed.rows[1].housing_type == "house" and parsed.rows[1].application_type == "general"


def test_parse_csv_cp949_and_missing_address_column() -> None:
    csv_data = "순번,소재지\n1,전주시 완산구 효자동2가 1243-3\n".encode("cp949")
    parsed = parse_upload("list.csv", csv_data)
    assert parsed.rows[0].address == "전주시 완산구 효자동2가 1243-3"

    with pytest.raises(HTTPException) as exc:
        parse_upload("bad.csv", "번호,이름\n1,홍길동\n".encode("utf-8"))
    assert "소재지 열" in exc.value.detail


def test_header_matches_by_containment_after_exact() -> None:
    # 15건.xlsx 처럼 「LH 파일 주소」로 적힌 열도 소재지로 잡는다. 「x」 같은 한 글자 키는 포함 검사에 안 쓴다.
    data = _xlsx([["접수번호", "LH 파일 주소", "xlsx비고"], ["002", "전주시 완산구 효자동2가 1243-3", "메모"]])
    parsed = parse_upload("15건.xlsx", data)
    assert parsed.columns["address"] == "LH 파일 주소"
    assert "lng" not in parsed.columns
    assert parsed.rows[0].address == "전주시 완산구 효자동2가 1243-3"


def test_parse_lh_template_columns_extras_and_warnings() -> None:
    # LH 「심사지 리스트 양식」(09/22): 접수번호가 순번보다 식별자로 앞서고, 분류·신청유형이 행별 유형이 된다.
    data = _xlsx(
        [
            ["순번", "차수", "접수번호", "접수일자", "매도자명", "분류", "신청유형", "지역", "상세주소", "경도", "위도"],
            [1, "1001", "001", "26.04.13.", "홍길동", "주거용 오피스텔", "청년", "전주시", "완산구 효자동 1243-2", 127.10403, 35.815217],
            [5, "1001", "005", "26.04.30.", "", "도시형생활주택 + 주거용오피스텔", "신혼", "익산시", "창인동1가 19 외 2필지", None, None],
            [6, "1001", "006", "26.04.30.", "", "도시형생활주택", "복지", "익산시", "부송동 764-10, -7", None, None],
        ]
    )
    parsed = parse_upload("양식.xlsx", data)
    rows = {r.id: r for r in parsed.rows}

    assert list(rows) == ["001", "005", "006"]
    assert parsed.columns["housing"] == "분류" and parsed.columns["차수"] == "차수"
    assert rows["001"].housing_type == "officetel" and rows["001"].application_type == "youth"
    assert rows["001"].extras == {"차수": "1001", "접수일자": "26.04.13.", "매도자명": "홍길동"}
    assert rows["001"].warnings == []
    # 혼합 분류는 정하지 않고 담당자가 고르게 한다. 「외 N필지」만 있으면 대표 필지만 판정됨을 알린다.
    assert rows["005"].housing_type is None
    assert any("양쪽" in w for w in rows["005"].warnings)
    assert any("외 N필지" in w for w in rows["005"].warnings)
    # 도시형생활주택은 주택. 모르는 신청유형은 경고로 남긴다.
    assert rows["006"].housing_type == "house" and rows["006"].application_type is None
    assert any("신청유형" in w for w in rows["006"].warnings)
    assert not any("외 N필지" in w for w in rows["006"].warnings)


def test_sheet_name_fills_blank_type_columns() -> None:
    # LH 양식 파일은 시트 이름이 「청년」이고 분류·신청유형 칸이 비어 있다. 시트 이름을 빈 칸의 기본값으로 쓴다.
    book = Workbook()
    sheet = book.active
    sheet.title = "청년"
    sheet.append(["접수번호", "분류", "신청유형", "지역", "상세주소"])
    sheet.append(["001", "", "", "전주시", "완산구 효자동2가 1243-3"])
    sheet.append(["002", "주택", "일반", "전주시", "덕진구 금암동 473-6"])
    buffer = io.BytesIO()
    book.save(buffer)
    rows = parse_upload("청년.xlsx", buffer.getvalue()).rows

    assert rows[0].application_type == "youth" and rows[0].housing_type is None
    assert rows[0].application_label.startswith("시트")
    # 칸에 값이 있으면 시트 이름보다 칸이 앞선다.
    assert rows[1].application_type == "general" and rows[1].housing_type == "house"


def test_template_download_has_fixed_header() -> None:
    from openpyxl import load_workbook

    from app.screening.batch import TEMPLATE_COLUMNS, build_template

    sheet = load_workbook(io.BytesIO(build_template())).active
    assert [c.value for c in sheet[1]] == list(TEMPLATE_COLUMNS)
    assert parse_upload("t.xlsx", build_template()).rows[0].application_type == "youth"


def test_type_labels_map_loosely() -> None:
    assert housing_type_from_label("주거용 오피스텔") == "officetel"
    assert housing_type_from_label("아파트") == "house"
    assert housing_type_from_label("도시형생활주택") == "house"
    assert housing_type_from_label("도시형생활주택 + 주거용오피스텔") is None
    assert housing_type_from_label("") is None
    assert application_type_from_label("신혼·신생아") == "newlywed"
    assert application_type_from_label("GENERAL") == "general"
    assert application_type_from_label("모름") is None


class FakeKakao:
    def __init__(self) -> None:
        self.calls: list[str] = []

    async def geocode(self, query: str) -> list[GeocodeCandidate]:
        self.calls.append(query)
        if "없는주소" in query:
            return []
        return [GeocodeCandidate(id="k1", name=query, address=query, coordinates=CENTER)]


class FakeResolver:
    async def resolve(self, request):  # noqa: ANN001
        return HazardParcelResolveResponse(
            parcels=[HazardParcel(parcel_id="p1", address=request.address)],
            note="대표 필지",
        )


class SpyScreening(ScreeningService):
    def __init__(self) -> None:
        super().__init__(
            hazard=FakeHazardService([category("a", "no_conflict_in_snapshot")]),  # type: ignore[arg-type]
            amenities=FakeAmenityCollector(collections()),  # type: ignore[arg-type]
            demo_mode=False,
        )
        self.requests: list[ScreeningRequest] = []

    async def screen(self, request, progress=None):  # noqa: ANN001
        self.requests.append(request)
        return await super().screen(request, progress)


def _batch(rows: list[BatchRowInput]) -> BatchStatus:
    return BatchStatus(
        batch_id="b1",
        status="queued",
        created_at=datetime.now(UTC),
        housing_type="house",
        application_type="general",
        total=len(rows),
        rows=[
            BatchRowStatus(
                id=r.id, address=r.address,
                housing_type=r.housing_type or "house",
                application_type=r.application_type or "general",
            )
            for r in rows
        ],
    )


@pytest.mark.asyncio
async def test_run_batch_geocodes_only_rows_without_coordinates_and_isolates_failures() -> None:
    rows = [
        BatchRowInput(id="1", address="전주시 A", lat=CENTER.lat, lng=CENTER.lng),
        BatchRowInput(id="2", address="전주시 B", application_type="youth"),
        BatchRowInput(id="3", address="없는주소"),
    ]
    batch = _batch(rows)
    screening, kakao = SpyScreening(), FakeKakao()

    await run_batch(
        batch, rows, screening=screening, resolver=FakeResolver(), kakao=kakao,  # type: ignore[arg-type]
        rule_pack_id="", cancel=asyncio.Event(), concurrency=1,
    )

    assert batch.status == "completed" and batch.done == 3
    assert kakao.calls == ["전주시 B", "없는주소"]
    by_id = {r.id: r for r in batch.rows}
    assert by_id["1"].status == "completed" and by_id["1"].living_score is not None
    # 행별 유형이 일괄 기본값을 이긴다.
    assert screening.requests[1].site.application_type == "youth"
    assert by_id["3"].status == "failed" and "주소를 찾지 못했" in by_id["3"].error

    sheets = export_rows(batch)
    assert sheets["종합요약"][0][0] == "접수번호"
    assert len(sheets["종합요약"]) == 4
    assert sheets["종합요약"][3][-1].startswith("주소를 찾지 못했")
    # 완료 건은 1차 항목·2차 요건 행이 붙는다.
    assert any(r[0] == "1" for r in sheets["1차상세"][1:])
    assert any(r[0] == "1" and r[2] == "대중교통 접근성" for r in sheets["2차상세"][1:])


@pytest.mark.asyncio
async def test_cancel_skips_remaining_rows() -> None:
    rows = [BatchRowInput(id=str(i), address=f"전주시 {i}", lat=CENTER.lat, lng=CENTER.lng) for i in range(3)]
    batch = _batch(rows)
    cancel = asyncio.Event()
    cancel.set()

    await run_batch(
        batch, rows, screening=SpyScreening(), resolver=FakeResolver(), kakao=FakeKakao(),  # type: ignore[arg-type]
        rule_pack_id="", cancel=cancel,
    )

    assert batch.status == "cancelled"
    assert all(r.status == "cancelled" for r in batch.rows)
