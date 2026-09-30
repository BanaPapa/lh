"""결과 Excel — 4시트(종합요약·1차 상세·2차 상세·데이터 스냅샷)와 담당자 판단 반영."""

from __future__ import annotations

import io
import os
from datetime import UTC, datetime

os.environ.setdefault("DEMO_MODE", "true")
os.environ.setdefault("KAKAO_REST_API_KEY", "")

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.screening.batch import BatchRowStatus, BatchStatus, batch_results, batches
from app.screening.export_xlsx import ExportSite, ItemJudgement, build_workbook
from app.screening.models import ScreeningRequest, ScreeningResult
from app.services.facility_store import SyncState
from app.screening.service import ScreeningService

from tests.test_screening_scorebook import (  # noqa: E402
    TWO_TRANSIT,
    FakeAmenityCollector,
    FakeHazardService,
    category,
    collections,
    site,
)

SHEETS = ["종합요약", "1차 상세", "2차 상세", "데이터 스냅샷"]


class FakeStore:
    def sync_states(self) -> list[SyncState]:
        return [
            SyncState("lodgings", datetime(2026, 9, 1, tzinfo=UTC), 10),
            SyncState("gas_station", datetime(2026, 9, 20, tzinfo=UTC), 5),
            SyncState("empty", datetime(2026, 9, 25, tzinfo=UTC), 0),
        ]


async def _result() -> ScreeningResult:
    hazard = FakeHazardService([
        category("gas", "exclusion_match", label="주유소", nearest=20),
        category("lpg", "no_conflict_in_snapshot", label="LPG 충전소"),
    ])
    hazard.facility_store = FakeStore()  # type: ignore[attr-defined]
    service = ScreeningService(
        hazard=hazard,  # type: ignore[arg-type]
        amenities=FakeAmenityCollector(collections(**TWO_TRANSIT)),  # type: ignore[arg-type]
    )
    return await service.screen(ScreeningRequest(site=site()))


def _rows(data: bytes, sheet: str) -> list[tuple]:
    book = load_workbook(io.BytesIO(data))
    return [tuple(r) for r in book[sheet].iter_rows(values_only=True)]


@pytest.mark.asyncio
async def test_result_carries_data_snapshot() -> None:
    result = await _result()
    snapshot = result.data_snapshot
    assert snapshot is not None
    # 레코드 0건 데이터셋은 판에서 뺀다.
    assert snapshot.localdata_dataset_count == 2
    assert snapshot.localdata_record_count == 15
    assert snapshot.localdata_synced_at_min == datetime(2026, 9, 1, tzinfo=UTC)
    assert snapshot.localdata_synced_at_max == datetime(2026, 9, 20, tzinfo=UTC)
    assert {o.key for o in snapshot.options} >= {"bus_headway_filter"}


@pytest.mark.asyncio
async def test_workbook_has_four_sheets_with_judgements() -> None:
    result = await _result()
    data = build_workbook([
        ExportSite(
            receipt_no="001",
            result=result,
            judgements={"gas": ItemJudgement(state="false_positive", memo="현장 확인 폐업")},
        ),
        ExportSite(receipt_no="002", address="없는주소", error="주소를 찾지 못했습니다."),
    ])

    assert load_workbook(io.BytesIO(data)).sheetnames == SHEETS

    summary = _rows(data, "종합요약")
    assert summary[0][:4] == ("접수번호", "소재지", "분류", "신청유형")
    assert summary[1][0] == "001" and summary[1][4] == result.stage_one.verdict_label
    assert summary[1][13] == "1/2"  # 담당자 확인 1건/2항목
    assert summary[2][0] == "002" and summary[2][-1] == "주소를 찾지 못했습니다."

    one = _rows(data, "1차 상세")
    head = one[0]
    gas = next(r for r in one[1:] if r[3] == "주유소")
    assert gas[head.index("담당자 판단")] == "오탐"
    assert gas[head.index("메모")] == "현장 확인 폐업"
    lpg = next(r for r in one[1:] if r[3] == "LPG 충전소")
    assert lpg[head.index("담당자 판단")] == "미확인"

    two = _rows(data, "2차 상세")
    assert two[0][5] == "시설군" and two[0][7] == "시설명"
    assert any(r[2] == "대중교통 접근성" for r in two[1:])

    snap = dict((r[0], r[1]) for r in _rows(data, "데이터 스냅샷") if r and r[0])
    assert "생성 시각" in snap
    assert snap["규칙팩 버전"].startswith(result.rule_pack_id)
    assert snap["인허가(LOCALDATA) 캐시 동기화"].startswith("2026-09-01")
    assert any(k.startswith("적용 판정 옵션") for k in snap)


@pytest.mark.asyncio
async def test_single_export_endpoint_returns_xlsx() -> None:
    result = await _result()
    client = TestClient(app)
    response = client.post(
        "/api/screening/export.xlsx",
        json={"sites": [{
            "receipt_no": "A-1",
            "result": result.model_dump(mode="json"),
            "judgements": {"gas": {"state": "confirmed", "memo": ""}},
        }]},
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "filename*=UTF-8''" in response.headers["content-disposition"]
    one = _rows(response.content, "1차 상세")
    gas = next(r for r in one[1:] if r[3] == "주유소")
    assert gas[one[0].index("담당자 판단")] == "확인 완료"


@pytest.mark.asyncio
async def test_batch_legacy_export_endpoint_uses_server_results_and_judgements() -> None:
    """옛 4시트 묶음은 export.legacy.xlsx 로 남는다. 담당자 판단은 요청에서 받는다."""

    result = await _result()
    batch = BatchStatus(
        batch_id="bx", status="completed", created_at=datetime.now(UTC),
        housing_type="house", application_type="general", total=2,
        rows=[
            BatchRowStatus(id="1", address="전주시 A", housing_type="house", application_type="general",
                           status="completed", screening_id=result.screening_id),
            BatchRowStatus(id="2", address="없는주소", housing_type="house", application_type="general",
                           status="failed", error="주소를 찾지 못했습니다.", message="실패"),
        ],
    )
    batches["bx"] = batch
    batch_results[("bx", "1")] = result
    try:
        client = TestClient(app)
        response = client.post(
            "/api/screening/batch/bx/export.legacy.xlsx",
            json={"judgements": {result.screening_id: {"gas": {"state": "evidence_requested", "memo": "공문 요청"}}}},
        )
        assert response.status_code == 200
        summary = _rows(response.content, "종합요약")
        assert [r[0] for r in summary[1:]] == ["1", "2"]
        assert summary[2][-1] == "주소를 찾지 못했습니다."
        one = _rows(response.content, "1차 상세")
        gas = next(r for r in one[1:] if r[3] == "주유소")
        assert gas[0] == "1"
        assert gas[one[0].index("담당자 판단")] == "추가 증빙 요청"
        assert gas[one[0].index("메모")] == "공문 요청"

        assert client.post("/api/screening/batch/nope/export.legacy.xlsx", json={}).status_code == 404
    finally:
        batches.pop("bx", None)
        batch_results.pop(("bx", "1"), None)


# ---------------------------------------------------------------------------
# 일괄 심사 내려받기 — 결과표 + 건별 심사표(Excel·PDF)
# ---------------------------------------------------------------------------
def _batch(result: ScreeningResult, batch_id: str = "bp") -> BatchStatus:
    """완료 2건(하나는 차수·매도자 있음) + 실패 1건. 화면 결과표와 같은 모양이다."""

    from app.screening.batch import BatchCriterionScore, _fill_row

    done = BatchRowStatus(
        id="001", address="전주시 A", housing_type="house", application_type="youth",
        extras={"차수": "1", "매도자명": "홍길동"}, status="completed", message=result.verdict_label,
    )
    _fill_row(done, result, parcel_count=2, note="필지 2건 합집합")
    done.site_address = "전북특별자치도 전주시 A"
    partial = BatchRowStatus(
        id="002", address="전주시 B", housing_type="officetel", application_type="general",
        status="completed", message="적격", screening_id="other", verdict="pass", verdict_label="적격",
        living_score=None, living_score_min=20, living_score_max=28, living_maximum=40, determined=False,
        criteria=[BatchCriterionScore(key="transit", label="대중교통 접근성", awarded=None, awarded_min=12, awarded_max=20, maximum=20)],
    )
    failed = BatchRowStatus(
        id="003", address="없는주소", housing_type="house", application_type="general",
        status="failed", error="주소를 찾지 못했습니다.", message="실패",
    )
    return BatchStatus(
        batch_id=batch_id, status="completed", created_at=datetime.now(UTC),
        housing_type="house", application_type="general", total=3, done=3, progress=100,
        rows=[done, partial, failed],
    )


SCREEN_COLUMNS = ["접수번호", "차수·접수일자·매도자", "소재지", "분류 · 신청유형", "1차 판정", "2차 점수",
                  "교통", "주거", "교육", "가점", "필지", "상태"]


@pytest.mark.asyncio
async def test_batch_report_table_matches_screen() -> None:
    from app.screening.batch import BatchExportRequest, build_report

    result = await _result()
    batch = _batch(result)
    batch_results[("bp", "001")] = result
    try:
        report = build_report(batch, BatchExportRequest(), datetime.now(UTC))
    finally:
        batch_results.pop(("bp", "001"), None)

    assert report.columns == SCREEN_COLUMNS
    assert [r[0] for r in report.rows] == ["001", "002", "003"]
    first, second, third = report.rows
    assert first[1] == "1 · 홍길동"
    assert first[2] == "전주시 A\n대표 필지 전북특별자치도 전주시 A"
    assert first[3] == "주택 · 청년"
    assert first[4] == result.verdict_label
    assert first[5] == f"{result.stage_two.living_score}/{result.stage_two.living_maximum}"
    assert first[10] == "2" and first[11] == result.verdict_label
    # 미확정 건은 화면처럼 범위로, 없는 항목은 「·」.
    assert second[5] == "20~28/40" and second[6] == "12~20" and second[7] == "·"
    # 실패 건은 오류를 소재지 아래 줄에, 점수는 「—」.
    assert third[2] == "없는주소\n주소를 찾지 못했습니다." and third[4] == "—" and third[11] == "실패"
    assert len(report.sites) == 3
    assert report.sites[2].blocks[-1].text.startswith("심사 결과 없음 — 주소를 찾지 못했")


@pytest.mark.asyncio
async def test_site_report_carries_sheet_contents() -> None:
    from app.screening.export_report import site_report

    result = await _result()
    site = ExportSite(
        receipt_no="001", result=result, extras={"차수": "1"}, resolve_note="필지 2건 합집합",
        judgements={"gas": ItemJudgement(state="false_positive", memo="현장 확인 폐업")},
    )
    report = site_report(site, 1)
    blocks = report.blocks
    kinds = [b.kind for b in blocks]
    assert kinds[0] == "kv"
    info = dict((k, v) for k, v in blocks[0].rows)
    assert info["차수"] == "1" and info["필지 비고"] == "필지 2건 합집합"
    assert info["데이터 판"].startswith("데이터 판 인허가(LOCALDATA) 캐시 2026-09-01~2026-09-20")
    headings = [b.text for b in blocks if b.kind == "heading"]
    assert any(h.startswith("종합 판정 — ") for h in headings)
    assert "1차 매입제외 판정 — 주거환경 저해시설" in headings
    assert any(h.startswith("2차 생활편의성 배점 — ") for h in headings)
    assert "배점 근거" in headings and "담당자 판단 · 메모" in headings

    one = next(b for b in blocks if b.kind == "table" and b.columns[0] == "항목")
    assert one.columns == ["항목", "기준거리", "최근접", "결과", "근거", "원천 · 데이터", "측정 기준", "담당자 판단"]
    gas = next(r for r in one.rows if r[0].startswith("주유소"))
    assert gas[1] == "50m" and gas[2] == "20.0m" and gas[3] == "매입제외"
    assert gas[7] == "오탐 — 메모: 현장 확인 폐업"
    assert one.group_rows, "규칙별 묶음 머리행이 있어야 한다"

    reasons = next(b for b in blocks if b.kind == "table" and b.columns[0] == "평가항목")
    transit_criterion = result.stage_two.criteria[0]
    transit = next(r for r in reasons.rows if r[0] == transit_criterion.label)
    assert transit[1] == f"{transit_criterion.awarded} / {transit_criterion.maximum}"
    groups = [b for b in blocks if b.kind == "table" and b.columns[0] == "시설군"]
    assert groups, "시설군 연결상세 표가 있어야 한다"
    subway = next(r for b in groups for r in b.rows if r[0].startswith("지하철역"))
    assert subway[1] == "1건" and subway[2] == "400.0m"
    hits = [r for b in groups for i, r in enumerate(b.rows) if i in b.sub_rows]
    assert any(r[0].startswith("└ subway-0") for r in hits)

    # PDF 용 압축은 시설 목록을 시설군 칸 한 줄로 모은다.
    compact = site_report(site, 1, compact=True)
    compact_groups = [b for b in compact.blocks if b.kind == "table" and b.columns[0] == "시설군"]
    assert not any(b.sub_rows for b in compact_groups)
    assert any("반영 1곳: subway-0 400.0m" in r[4] for b in compact_groups for r in b.rows)


@pytest.mark.asyncio
async def test_batch_workbook_has_table_sheet_then_site_sheets() -> None:
    from app.screening.batch import BatchExportRequest, build_report
    from app.screening.export_report import build_batch_workbook

    result = await _result()
    batch = _batch(result)
    batch_results[("bp", "001")] = result
    try:
        data = build_batch_workbook(build_report(batch, BatchExportRequest(), datetime.now(UTC)))
    finally:
        batch_results.pop(("bp", "001"), None)
    book = load_workbook(io.BytesIO(data))
    assert book.sheetnames == ["결과표", "1. 001", "2. 002", "3. 003"]
    table = [tuple(r) for r in book["결과표"].iter_rows(values_only=True)]
    header_index = next(i for i, r in enumerate(table) if r[0] == "접수번호")
    assert list(table[header_index]) == SCREEN_COLUMNS
    assert [r[0] for r in table[header_index + 1:]] == ["001", "002", "003"]
    detail = [r[0] for r in book["1. 001"].iter_rows(values_only=True) if r[0]]
    assert detail[0].startswith("1. 접수번호 001 · ")
    assert "1차 매입제외 판정 — 주거환경 저해시설" in detail
    assert "항목" in detail and "시설군" in detail
    assert any(str(v).startswith("심사 결과 없음") for v in (r[0] for r in book["3. 003"].iter_rows(values_only=True)) if v)


@pytest.mark.asyncio
async def test_batch_pdf_is_a4_portrait_with_korean_font_and_page_per_site() -> None:
    from pypdf import PdfReader

    from app.screening.batch import BatchExportRequest, build_report
    from app.screening.export_pdf import FONT_NAME, build_batch_pdf, register_fonts

    register_fonts()
    from reportlab.pdfbase import pdfmetrics

    assert FONT_NAME in pdfmetrics.getRegisteredFontNames()

    result = await _result()
    batch = _batch(result)
    batch_results[("bp", "001")] = result
    try:
        data = build_batch_pdf(build_report(batch, BatchExportRequest(), datetime.now(UTC), compact=True))
    finally:
        batch_results.pop(("bp", "001"), None)
    reader = PdfReader(io.BytesIO(data))
    assert len(reader.pages) >= 1 + len(batch.rows)
    first = reader.pages[0]
    assert float(first.mediabox.width) < float(first.mediabox.height)  # A4 세로
    assert round(float(first.mediabox.width)) == 595 and round(float(first.mediabox.height)) == 842
    text = first.extract_text()
    assert "일괄 심사 결과" in text and "접수번호" in text and "전주시 A" in text
    fonts = {
        str(font.get_object().get("/BaseFont"))
        for page in reader.pages
        for font in (page["/Resources"].get("/Font") or {}).values()
    }
    assert any("NanumGothic" in name for name in fonts)
    # 2쪽부터 건별 심사표 — 순서대로.
    assert "접수번호 001" in reader.pages[1].extract_text()
    assert "1차 매입제외 판정" in reader.pages[1].extract_text()
    remaining = "".join(page.extract_text() for page in reader.pages[2:])
    assert "접수번호 002" in remaining and "접수번호 003" in remaining and "심사 결과 없음" in remaining


@pytest.mark.asyncio
async def test_batch_export_endpoints_return_xlsx_and_pdf() -> None:
    result = await _result()
    batch = _batch(result, "bq")
    batches["bq"] = batch
    batch_results[("bq", "001")] = result
    try:
        client = TestClient(app)
        xlsx = client.post("/api/screening/batch/bq/export.xlsx", json={"judgements": {}})
        assert xlsx.status_code == 200
        assert xlsx.headers["content-type"].startswith("application/vnd.openxmlformats")
        assert xlsx.headers["content-disposition"].startswith("attachment; filename=LH_batch_")
        assert "filename*=UTF-8''LH_%EC%9D%BC%EA%B4%84%EC%8B%AC%EC%82%AC_" in xlsx.headers["content-disposition"]
        assert load_workbook(io.BytesIO(xlsx.content)).sheetnames[0] == "결과표"

        pdf = client.post("/api/screening/batch/bq/export.pdf", json={"judgements": {}})
        assert pdf.status_code == 200
        assert pdf.headers["content-type"] == "application/pdf"
        assert pdf.content.startswith(b"%PDF-") and pdf.headers["content-disposition"].endswith(".pdf")

        legacy = client.post("/api/screening/batch/bq/export.legacy.xlsx", json={})
        assert legacy.status_code == 200 and load_workbook(io.BytesIO(legacy.content)).sheetnames == SHEETS
        assert client.post("/api/screening/batch/nope/export.pdf", json={}).status_code == 404
    finally:
        batches.pop("bq", None)
        batch_results.pop(("bq", "001"), None)


@pytest.mark.asyncio
async def test_snapshot_export_builds_files_without_server_state() -> None:
    """서버가 새로 켜져 결과가 사라져도, 브라우저가 보관한 결과(원본·경계 좌표 제외)로 같은 파일을 만든다."""

    result = await _result()
    batch = _batch(result, "gone")
    trimmed = result.model_dump(mode="json")
    trimmed["hazard_review"] = None
    for parcel in trimmed["site"]["parcels"]:
        parcel["geometry"] = []
    for item in trimmed["stage_one"]["items"]:
        for facility in item["facilities"]:
            facility["geometry"] = []
    for criterion in trimmed["stage_two"]["criteria"]:
        for group in criterion["groups"]:
            for hit in group["hits"]:
                hit["facility_ring"] = []
    payload = {
        "batch": batch.model_dump(mode="json"),
        "results": {"001": trimmed},
        "judgements": {result.screening_id: {"gas": {"state": "confirmed", "memo": ""}}},
    }
    client = TestClient(app)
    assert "gone" not in batches
    xlsx = client.post("/api/screening/batch/export.xlsx", json=payload)
    assert xlsx.status_code == 200, xlsx.text
    book = load_workbook(io.BytesIO(xlsx.content))
    assert book.sheetnames == ["결과표", "1. 001", "2. 002", "3. 003"]
    detail = [r for r in book["1. 001"].iter_rows(values_only=True) if r[0] == "주유소" and r[1] == "50m"]
    assert detail and detail[0][7] == "확인 완료"
    # 결과를 못 받아 둔 건은 그 사실을 적는다.
    missing = [r[0] for r in book["2. 002"].iter_rows(values_only=True) if r[0]]
    assert any(str(v).startswith("심사 결과 없음 — 브라우저에 보관된 결과 없음") for v in missing)
    pdf = client.post("/api/screening/batch/export.pdf", json=payload)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF-")


def test_finished_batches_are_evicted_oldest_first() -> None:
    from datetime import timedelta

    from app.screening.batch import evict_finished_batches

    base = datetime.now(UTC)
    ids = [f"ev{i}" for i in range(4)]
    try:
        for index, batch_id in enumerate(ids):
            status = "running" if index == 0 else "completed"
            batches[batch_id] = BatchStatus(
                batch_id=batch_id, status=status, created_at=base + timedelta(minutes=index),
                housing_type="house", application_type="general", total=1,
                rows=[BatchRowStatus(id="1", address="a", housing_type="house", application_type="general")],
            )
            batch_results[(batch_id, "1")] = None  # type: ignore[assignment]
        removed = evict_finished_batches(keep=1)
        # 진행 중(ev0)은 남고, 끝난 것 중 오래된 ev1·ev2 만 지운다.
        assert removed == ["ev1", "ev2"]
        assert "ev0" in batches and "ev3" in batches
        assert ("ev1", "1") not in batch_results and ("ev3", "1") in batch_results
    finally:
        for batch_id in ids:
            batches.pop(batch_id, None)
            batch_results.pop((batch_id, "1"), None)

