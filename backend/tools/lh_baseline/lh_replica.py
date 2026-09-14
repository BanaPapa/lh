"""LH 기존 웹앱 2종의 판정 로직을 그대로 옮긴 재현본.

03_ 1단계 유해시설 모델 · 04_ 2단계 생활입지평가 모델의 JS 를 1:1로 포팅했다.
숫자를 바꾸지 않는다. 대조의 기준선이므로 '고쳐서' 옮기면 대조가 성립하지 않는다.
"""

from __future__ import annotations

import glob
import math
import os
import unicodedata

import openpyxl

# LH 원본 자료(01_ 신청 원장 · 02_ 유해시설 · 00_ 시군별 시설 마스터)가 있는 폴더.
# 저장소에 담지 않는다 — 신청 원장에 개인정보가 섞여 있다.
# 환경변수 LH_SOURCE_DIR 로 지정한다.
#
# 경로는 임포트 시점에 굳으므로, 그 전에 `.env` 를 프로세스 환경변수로 올려 둔다.
# 호출자가 나중에 올려 주기를 기다리면 BASE 가 빈 문자열로 굳어 원장을 못 찾는다.
from app.settings_api import store as _settings_store

_settings_store.hydrate_process_env()

BASE = os.environ.get("LH_SOURCE_DIR", "")

# 01_ 신청 원장 경로. 스냅샷 메타에 어떤 원장을 읽었는지 남기려고 상수로 둔다.
LEDGER_PATH = os.path.join(BASE, "01_ 매입약정 신청자 전체파일(260430).xlsx")


# ---------------------------------------------------------------------------
# 03_ 1단계 — haversineMeters / getThresholds / analyzeSites
# ---------------------------------------------------------------------------
EARTH_R = 6_371_000.0  # 03_ 의 EARTH_R


def haversine_m(lat1: float, lon1: float, lat2: float, lon2: float) -> float:
    r = math.pi / 180
    la1, lo1, la2, lo2 = lat1 * r, lon1 * r, lat2 * r, lon2 * r
    dlat, dlon = la2 - la1, lo2 - lo1
    a = math.sin(dlat / 2) ** 2 + math.cos(la1) * math.cos(la2) * math.sin(dlon / 2) ** 2
    return EARTH_R * 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))


def thresholds(major: str, hate_radius: int = 1000) -> tuple[int, int]:
    """03_ getThresholds() 그대로. 알 수 없는 분류는 보수적으로 25/50."""

    if major in ("숙박시설", "위락시설", "위험시설"):
        return 25, 50
    if major == "혐오시설":
        return 500, hate_radius
    return 25, 50


def analyze_site(lat: float, lon: float, hazards: list[dict], hate_radius: int = 1000) -> dict:
    """03_ analyzeSites() 의 사업지 1건 처리 부분 그대로."""

    worst_rank = 0
    worst = None
    worst_d: float | None = None
    nearest = None
    nearest_d = math.inf

    for hz in hazards:
        d = haversine_m(lat, lon, hz["lat"], hz["lon"])
        if d < nearest_d:
            nearest_d, nearest = d, hz
        ex, ca = thresholds(hz["major"], hate_radius)
        rank = 2 if d <= ex else (1 if d <= ca else 0)
        # 원본의 동점 처리: 같은 rank 면 더 가까운 시설로 교체
        if rank > worst_rank or (
            rank == worst_rank
            and rank > 0
            and d < (math.inf if worst_d is None else worst_d)
        ):
            worst_rank, worst, worst_d = rank, hz, d

    if worst_rank == 0:
        return {
            "status": "PASS",
            "nearest": nearest,
            "distance": nearest_d if nearest else None,
        }
    return {
        "status": "EXCLUDE" if worst_rank == 2 else "CAUTION",
        "nearest": worst,
        "distance": worst_d,
    }


# ---------------------------------------------------------------------------
# 04_ 2단계 — hav / calcScores  (일반형 15·15·10 하드코딩 그대로)
# ---------------------------------------------------------------------------
def hav_km(la1: float, lo1: float, la2: float, lo2: float) -> float:
    r = math.pi / 180
    dlat, dlng = (la2 - la1) * r, (lo2 - lo1) * r
    a = (
        math.sin(dlat / 2) ** 2
        + math.cos(la1 * r) * math.cos(la2 * r) * math.sin(dlng / 2) ** 2
    )
    return 6371.0 * 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))


def calc_scores(rows: list[dict]) -> tuple[int, int, int]:
    """04_ calcScores() 그대로. rows 는 dist(km) 가 채워진 시설 목록."""

    trans = [r for r in rows if r["big"] == "대중교통"]
    t05 = sum(1 for r in trans if r["dist"] <= 0.5)
    t10 = sum(1 for r in trans if r["dist"] <= 1.0)
    t15 = sum(1 for r in trans if r["dist"] <= 1.5)
    s1 = 15 if t05 >= 2 else 12 if t05 >= 1 else 9 if t10 >= 1 else 6 if t15 >= 1 else 3

    def ex(cat: str, km: float) -> bool:
        return any(r["big"] == cat and r["dist"] <= km for r in rows)

    def sub(km: float) -> int:
        return len(
            {
                r["big"]
                for r in rows
                if r["dist"] <= km and r["big"] in ("공원", "문화시설", "공공시설")
            }
        )

    comm2, medi2 = ex("상업시설", 2), ex("의료시설", 2)
    comm3, medi3 = ex("상업시설", 3), ex("의료시설", 3)
    s2 = (
        15 if (comm2 and medi2 and sub(2) >= 2)
        else 12 if (comm3 and medi3 and sub(3) >= 1)
        else 9 if ((comm3 or medi3) and sub(3) >= 2)
        else 6 if ((comm3 or medi3) and sub(3) >= 1)
        else 3
    )

    def hs(kind: str, km: float) -> bool:
        return any(
            r["big"] == "학교" and kind in r["mid"] and r["dist"] <= km for r in rows
        )

    s3 = (
        10 if (hs("초등", 0.5) and hs("중", 0.5) and hs("초등", 1) and hs("중", 1) and hs("고", 1))
        else 8 if (hs("초등", 1) and hs("중", 1.5) and hs("고", 1.5))
        else 6 if (hs("초등", 2) and hs("중", 2) and hs("고", 3))
        else 4 if (hs("초등", 3) and hs("중", 3))
        else 2
    )
    return s1, s2, s3


# ---------------------------------------------------------------------------
# 데이터 적재
# ---------------------------------------------------------------------------
def _norm(name: str) -> str:
    # 맥에서 만든 파일명은 자모 분리(NFD)라 윈도우 경로 비교가 어긋난다.
    return unicodedata.normalize("NFC", name)


def load_hazards() -> list[dict]:
    """02_ 유해시설(전북 전체). 03_ 의 파서와 같은 컬럼을 쓴다."""

    path = os.path.join(BASE, "02_ 1단계_ 유해시설(전북 전체)_260715.xlsx")
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = [str(c or "").strip() for c in next(it)]
    idx = {name: header.index(name) for name in ("대분류", "중분류", "시설명", "주소")}
    i_lng = next(k for k, v in enumerate(header) if "경도" in v)
    i_lat = next(k for k, v in enumerate(header) if "위도" in v)

    rows: list[dict] = []
    for r in it:
        try:
            lat, lng = float(r[i_lat]), float(r[i_lng])
        except (TypeError, ValueError):
            continue  # 03_ 도 좌표 없는 행은 버린다
        rows.append(
            {
                "major": str(r[idx["대분류"]] or "").strip(),
                "mid": str(r[idx["중분류"]] or "").strip(),
                "name": str(r[idx["시설명"]] or "").strip(),
                "address": str(r[idx["주소"]] or "").strip(),
                "lat": lat,
                "lon": lng,
            }
        )
    return rows


def load_amenities() -> dict[str, list[dict]]:
    """00_ 시군별 시설 DB 마스터. 04_ 는 파일 하나만 올려 쓰므로 시군별로 나눠 둔다."""

    out: dict[str, list[dict]] = {}
    folder = os.path.join(BASE, "00_ 전북 시군별 좌표(260415)")
    for path in glob.glob(os.path.join(folder, "*.xlsx")):
        city = _norm(os.path.basename(path)).replace("시설_DB마스터_", "").replace(".xlsx", "")
        ws = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [str(c or "").strip() for c in next(it)]
        i_big, i_mid, i_nm = (header.index(k) for k in ("대분류", "중분류", "시설명"))
        i_lat = next(k for k, v in enumerate(header) if "위도" in v)
        i_lng = next(k for k, v in enumerate(header) if "경도" in v)
        rows: list[dict] = []
        for r in it:
            if not r[i_nm]:
                continue
            try:
                lat, lng = float(r[i_lat]), float(r[i_lng])
            except (TypeError, ValueError):
                continue
            rows.append(
                {
                    "big": str(r[i_big] or "").strip(),
                    "mid": str(r[i_mid] or "").strip(),
                    "nm": str(r[i_nm] or "").strip(),
                    "lat": lat,
                    "lng": lng,
                }
            )
        out[city] = rows
    return out


def load_sites() -> list[dict]:
    """01_ 매입약정 신청자 전체파일."""

    ws = openpyxl.load_workbook(LEDGER_PATH, read_only=True, data_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = [str(c or "").strip() for c in next(it)]
    idx = {n: header.index(n) for n in ("순번", "접수번호", "지역", "상세주소", "경도", "위도")}
    sites: list[dict] = []
    for r in it:
        if r[idx["접수번호"]] is None:
            continue
        try:
            lat, lng = float(r[idx["위도"]]), float(r[idx["경도"]])
            valid = 33.0 <= lat <= 39.5 and 124.0 <= lng <= 132.0
        except (TypeError, ValueError):
            lat = lng = float("nan")
            valid = False
        sites.append(
            {
                "seq": r[idx["순번"]],
                "no": str(r[idx["접수번호"]]).strip(),
                "region": str(r[idx["지역"]] or "").strip(),
                "detail": str(r[idx["상세주소"]] or "").strip(),
                "lat": lat,
                "lng": lng,
                "valid": valid,
            }
        )
    return sites


STATUS_LABEL = {"PASS": "통과", "CAUTION": "주의", "EXCLUDE": "제외", "NOCOORD": "좌표없음"}
