"""Meeting evidence runner. Keeps supplied applications and datasets unchanged."""
from __future__ import annotations
import argparse, asyncio, hashlib, json, re, sys, time
from pathlib import Path
from datetime import datetime, timezone, timedelta
from bs4 import BeautifulSoup

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parents[3]
INBOX = Path(r'C:\Users\Space\Documents\카카오톡 받은 파일')
JYS = Path(r'C:\Users\Space\Downloads\LH 심사 지원 앱_v3\app')
OUT = ROOT / 'docs' / 'reports' / 'meeting_20260911'
REF = INBOX / 'PoC 결과물 검증 기록부_2609.html'
MEETING = INBOX / 'LH_회의록_용역2차보고_2026-09-11.html'

def now(): return datetime.now(timezone(timedelta(hours=9))).isoformat()
def readj(p): return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def save(p, obj):
    p = Path(p); p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding='utf-8')
def digest(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()

def inspect(args):
    soup = BeautifulSoup((MEETING if args.meeting else REF).read_text(encoding='utf8'), 'html.parser')
    for x in soup(['style','script']): x.decompose()
    if args.section:
        start=soup.find(id=args.section)
        if start is None: raise ValueError(args.section)
        print(start.get_text(' ',strip=True))
        for el in start.next_siblings:
            if getattr(el,'name',None) in ('h2','h3'): break
            if hasattr(el,'get_text'): print(el.get_text(' ',strip=True))
    else:
        for i,t in enumerate(soup.find_all('table')):
            rows=[[c.get_text(' ',strip=True) for c in r.find_all(['th','td'],recursive=False)] for r in t.find_all('tr')]
            print(json.dumps({'table':i,'rows':rows if args.all else rows[:3]},ensure_ascii=False))

def sheets(args):
    import openpyxl
    paths=[INBOX/'LH 용역/결과물/AI공모전/01_ 매입약정 신청자 전체파일(260430).xlsx']
    paths.extend(p for p in INBOX.rglob('*.xlsx') if '17건' in p.name or '15건' in p.name)
    for p in paths:
        w=openpyxl.load_workbook(p,read_only=True,data_only=True)
        print('FILE',str(p))
        for s in w:
            print('SHEET',s.title,s.max_row,s.max_column)
            # Only headers until positions of sensitive columns are known.
            if '신청자' in p.name:
                for i,row in enumerate(s.iter_rows(values_only=True),1):
                    if i<=3: print(i,json.dumps(row,ensure_ascii=False,default=str))
            else:
                for i in [4,5,6,7,27,28,29,30,45]:
                    if i<=s.max_row: print(i,json.dumps([c.value for c in s[i]],ensure_ascii=False,default=str))
        w.close()

def prepare(args):
    import openpyxl, subprocess
    ledger_path=INBOX/'LH 용역/결과물/AI공모전/01_ 매입약정 신청자 전체파일(260430).xlsx'
    w=openpyxl.load_workbook(ledger_path,read_only=True,data_only=True)
    ledger={}
    for row in w['좌표'].iter_rows(min_row=2,values_only=True):
        no=str(row[2] or '').strip().zfill(3)
        if not no.isdigit() or int(no)==0 or not row[7] or not row[8]: continue
        ledger[no]={'seq':no,'address':f'{row[5]} {row[6]}','lat':float(row[8]),'lng':float(row[7]),'housing_type':'house','application_type':'general','group':'118'}
    w.close()
    assert len(ledger)==118, len(ledger)
    p17=INBOX/'LH 용역/2026 매입약정 서류심사_생활편의성 35점 이상 17건_26.09.07.작업.xlsx'
    w=openpyxl.load_workbook(p17,read_only=True,data_only=True); s=w['Sheet1']; cases=[]
    amap={'청년':'youth','일반':'general','신혼1':'newlywed','신혼2':'newlywed','다자녀':'multi_child','고령자':'senior'}
    for nr,ar,br,sr in [(4,5,6,7),(27,28,29,30)]:
        for col in range(4,s.max_column+1):
            name=s.cell(nr,col).value
            if not name: continue
            seq,app=str(name).split('_',1); seq=seq.zfill(3)
            label=str(s.cell(br,col).value); ht='officetel' if '오피스텔' in label else 'house'
            base={**ledger[seq],'group':'17','case_name':name,'application_type':amap[app],'building_label':label,'lh_actual':s.cell(sr,col).value,'score_source':f'Sheet1!{s.cell(sr,col).coordinate}','score_sheet_address':s.cell(ar,col).value}
            # Mixed building is explicit: primary house and secondary officetel.
            for housing in (['house','officetel'] if label.startswith('도시형생활주택') else [ht]):
                cases.append({**base,'housing_type':housing})
    w.close()
    assert len({c['seq'] for c in cases})==17
    soup=BeautifulSoup(REF.read_text(encoding='utf8'),'html.parser')
    tables=[]
    for t in soup.find_all('table'):
        tables.append([[c.get_text(' ',strip=True) for c in r.find_all(['th','td'],recursive=False)] for r in t.find_all('tr')])
    pilot=[{**ledger[n],'group':'pilot'} for n in ['001','004','006','041','105','060','085','095','104','106','108']]
    save(OUT/'inputs.json',{'created_at':now(),'ledger':list(ledger.values()),'cases17':cases,'pilot':pilot,'reference_tables':tables})
    paths=[REF,MEETING,ledger_path,p17,JYS/'data/snapshot/facilities.meta.json',JYS/'data/snapshot/facilities.json',JYS/'src/lib/engine/screen.ts',ROOT/'backend/app/screening/service.py',ROOT/'backend/app/hazard_review/service.py']
    save(OUT/'manifest.json',{'created_at':now(),'git_commit':subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip(),'files':[{'name':str(p),'size':p.stat().st_size,'sha256':digest(p)} for p in paths],'jys_snapshot':readj(JYS/'data/snapshot/facilities.meta.json'),'credential_values_recorded':False})
    print(json.dumps({'ledger':len(ledger),'cases17':len(cases),'pilot':len(pilot),'out':str(OUT)},ensure_ascii=False))

def brief(app,data):
    if app=='ours':
        return {'verdict':data.get('verdict'),'score':data.get('stage_two',{}).get('living_score'),'axes':{x['key']:x.get('awarded') for x in data.get('stage_two',{}).get('criteria',[])},'parcels':len(data.get('site',{}).get('parcels',[])),'missing':[x['key'] for x in data.get('hazard_review',{}).get('categories',[]) if x.get('status')=='dataset_missing']}
    return {'verdict':data.get('summary',{}).get('first_stage_status'),'axes':[{k:x.get(k) for k in ('axis','score','status')} for x in data.get('stage2',[])],'parcels':len(data.get('site',{}).get('parcels',[]))}

async def run_async(args):
    import httpx
    inputs=readj(OUT/'inputs.json')
    if args.group=='all':
        uniq={}
        for c in inputs['cases17']+inputs['pilot']+inputs['ledger']:
            uniq.setdefault((c['seq'],c['housing_type'],c['application_type']),c)
        rows=list(uniq.values())
    else: rows=inputs[{'17':'cases17','118':'ledger','pilot':'pilot'}[args.group]]
    if args.ids: rows=[c for c in rows if c['seq'] in args.ids.split(',')]
    apps=['ours','jys'] if args.app=='both' else [args.app]
    async with httpx.AsyncClient(timeout=600,trust_env=False) as client:
        async def one(app,c):
            key=f"{c['seq']}_{c['housing_type']}_{c['application_type']}"
            path=OUT/'raw'/app/f'{key}.json'
            if path.exists() and not args.force:
                existing=readj(path)
                if 'result' in existing:
                    print('CACHED',app,key,flush=True); return
            start=time.monotonic(); record={'started_at':now(),'app':app,'case':c}
            try:
                if app=='jys':
                    req={'application_id':f'MEETING-{key}','application_type_code':{'general':'GENERAL','youth':'YOUTH','newlywed':'NEWLYWED','multi_child':'MULTI_CHILD','senior':'ELDERLY'}[c['application_type']],'building_class':'주거용 오피스텔' if c['housing_type']=='officetel' else '주택','rule_pack':'YOUTH_DORM' if c['application_type']=='youth' else 'ELDERLY' if c['application_type']=='senior' else 'COMMON','parcels':[{'parcel_id':key,'jibun':c['address'],'geometry':{'type':'Point','coordinates':[c['lng'],c['lat']]}}]}
                    record['request']=req
                    r=await client.post('http://127.0.0.1:3000/api/screening/review',json=req)
                    record['http']=r.status_code
                    data=r.json()
                    if r.status_code!=200 or 'summary' not in data: raise ValueError(str(data)[:350])
                else:
                    resolve={'name':key,'address':c['address'],'coordinates':{'lat':c['lat'],'lng':c['lng']}}
                    r=await client.post('http://127.0.0.1:8000/api/hazard-review/parcels/resolve',json=resolve)
                    resolved=r.json(); record['resolve_request']=resolve; record['resolved']=resolved
                    if r.status_code!=200 or not resolved.get('parcels'): raise ValueError(str(resolved)[:350])
                    req={'site':{'name':key,'address':c['address'],'coordinates':resolve['coordinates'],'housing_type':c['housing_type'],'application_type':c['application_type'],'parcels':resolved['parcels']},'include_stage_two_on_fail':True,'requested_by':'meeting-verification-20260911'}
                    record['request']=req
                    r=await client.post('http://127.0.0.1:8000/api/screening/jobs',json=req)
                    r.raise_for_status(); job=r.json()['job_id']; record['job_id']=job
                    while True:
                        r=await client.get(f'http://127.0.0.1:8000/api/screening/jobs/{job}'); j=r.json()
                        if j.get('status')=='completed': data=j['result']; break
                        if j.get('status') in ('failed','cancelled'): raise ValueError(str(j.get('error') or j.get('message'))[:350])
                        if time.monotonic()-start>900: raise TimeoutError('900-second per-case limit')
                        await asyncio.sleep(2)
                record.update(result=data,finished_at=now(),elapsed_seconds=round(time.monotonic()-start,2))
                save(path,record)
                b=brief(app,data)
                print(json.dumps({'app':app,'key':key,'seconds':record['elapsed_seconds'],'brief':b},ensure_ascii=False),flush=True)
            except Exception as exc:
                record.update(error_type=type(exc).__name__,error=str(exc)[:500],finished_at=now(),elapsed_seconds=round(time.monotonic()-start,2)); save(path,record)
                print(json.dumps({'app':app,'key':key,'error':record['error']},ensure_ascii=False),flush=True)
        async def work(app):
            semaphore=asyncio.Semaphore(3 if app=='ours' else 1)
            async def bounded(c):
                async with semaphore: await one(app,c)
            await asyncio.gather(*(bounded(c) for c in rows))
        await asyncio.gather(*(work(a) for a in apps))

def run(args): asyncio.run(run_async(args))

def audit(args):
    from collections import Counter
    inputs=readj(OUT/'inputs.json')
    for app in ['ours','jys']:
        records=[readj(p) for p in (OUT/'raw'/app).glob('*.json')]
        print(app,'saved',len(records),'errors',[(x['case']['seq'],x.get('error')) for x in records if 'result' not in x])
    for c in inputs['cases17']:
        key=f"{c['seq']}_{c['housing_type']}_{c['application_type']}"
        vals={}
        for app in ['ours','jys']:
            p=OUT/'raw'/app/f'{key}.json'
            if not p.exists(): continue
            d=readj(p).get('result',{}); vals[app]=brief(app,d)
            if app=='ours':
                vals[app]['notes']=[{'criterion':x['key'],'note':x.get('note'),'groups':[{k:g.get(k) for k in ['key','state','note','count']} for g in x.get('groups',[])]} for x in d.get('stage_two',{}).get('criteria',[]) if not x.get('determined')]
        print(json.dumps({'key':key,'lh':c['lh_actual'],**vals},ensure_ascii=False))

def control(args):
    import httpx
    with httpx.Client(timeout=120,trust_env=False) as client:
        for key in ['060_house_youth','095_officetel_newlywed','085_house_general','108_officetel_general','010_house_general','025_house_general','084_house_general']:
            own=readj(OUT/'raw/ours'/f'{key}.json'); local=readj(OUT/'raw/jys'/f'{key}.json')
            pnus=[p['pnu'] for p in own['result']['site']['parcels'] if p.get('pnu')]
            if not pnus: continue
            req={**local['request'],'parcels':[{'pnu':p} for p in pnus]}
            start=time.monotonic(); r=client.post('http://127.0.0.1:3000/api/screening/review',json=req)
            d=r.json();save(OUT/'raw/jys_same_pnu'/f'{key}.json',{'started_at':now(),'case':own['case'],'request':req,'result':d,'http':r.status_code,'elapsed_seconds':round(time.monotonic()-start,2),'condition':'우리 앱이 선택한 PNU만 진용성 앱에 명시적으로 전달. 시설 스냅샷은 변경하지 않음.'})
            print(key,json.dumps(brief('jys',d),ensure_ascii=False))

def normalize(app, record):
    c=record['case'];d=record.get('result',{}); out={'seq':c['seq'],'address':c['address'],'housing_type':c['housing_type'],'application_type':c['application_type'],'elapsed_seconds':record.get('elapsed_seconds'),'error':record.get('error')}
    if app=='ours':
        two=d.get('stage_two',{}); cats=d.get('hazard_review',{}).get('categories',[])
        out.update(verdict=d.get('verdict'),score=two.get('living_score'),score_min=two.get('living_score_min'),score_max=two.get('living_score_max'),axes={x['key']:x.get('awarded') for x in two.get('criteria',[])},pnus=[x.get('pnu') for x in d.get('site',{}).get('parcels',[])],missing=[x['key'] for x in cats if x.get('status')=='dataset_missing'],missing_labels=[x.get('label',x['key']) for x in cats if x.get('status')=='dataset_missing'],geometry_sources=[x.get('geometry_source') for x in d.get('site',{}).get('parcels',[])],reasons=d.get('stage_one',{}).get('reasons',[]),reference_only=two.get('reference_only'),uncertain=[x.get('note') for x in two.get('criteria',[]) if not x.get('determined')])
    else:
        scores=[x.get('score') for x in d.get('stage2',[])]; score=sum(scores) if len(scores)==3 and all(isinstance(x,(int,float)) for x in scores) else None
        out.update(verdict={'EXCLUSION_MATCH':'fail','NO_CONFLICT_IN_SNAPSHOT':'pass','REVIEW_REQUIRED':'review'}.get(d.get('summary',{}).get('first_stage_status'),d.get('summary',{}).get('first_stage_status')),score=score,score_min=score,score_max=score,axes={ {'TRANSPORT':'transit','LIVING':'living','EDUCATION':'education'}[x['axis']]:x.get('score') for x in d.get('stage2',[])},pnus=[x.get('pnu') for x in d.get('site',{}).get('parcels',[])],missing=[x.get('group') for x in d.get('stage1',[]) if x.get('status')=='DATASET_MISSING'],reasons=[f"{x['group']}: "+', '.join(f"{f['facility_name']} {f['distance_m_display']}" for f in x.get('facilities',[])) for x in d.get('stage1',[]) if x.get('status') in ['EXCLUSION_MATCH','REVIEW_REQUIRED']],reference_only=d.get('summary',{}).get('second_stage_status')=='COMPLETED_NOT_PURCHASABLE',uncertain=[])
    return out

def summarize(args):
    from collections import Counter
    inputs=readj(OUT/'inputs.json'); records={}
    for app in ['ours','jys']:
        records[app]={p.stem:normalize(app,readj(p)) for p in (OUT/'raw'/app).glob('*.json')}
    def paired(cases):
        out=[]
        for c in cases:
            key=f"{c['seq']}_{c['housing_type']}_{c['application_type']}"
            out.append({'key':key,'case':c,'ours':records['ours'].get(key),'jys':records['jys'].get(key)})
        return out
    primary=[];seen=set()
    for c in inputs['cases17']:
        if c['seq'] not in seen: primary.append(c);seen.add(c['seq'])
    groups={'118':paired(inputs['ledger']),'17':paired(primary),'17_variants':paired(inputs['cases17']),'5':paired(inputs['pilot'][:5]),'focus':paired(inputs['pilot'])}
    groups['15']=[x for x in groups['17'] if x['case']['seq'] not in ['111','115']]
    metrics={}
    for name,rows in groups.items():
        m={'n':len(rows)}
        for app in ['ours','jys']:
            rs=[x[app] for x in rows if x[app] and not x[app].get('error')]
            m[app]={'completed':len(rs),'verdicts':dict(Counter(r['verdict'] for r in rs)),'score_determined':sum(r['score'] is not None for r in rs),'with_missing_source':sum(bool(r['missing']) for r in rs),'parcel_count':dict(Counter(len(r['pnus']) for r in rs))}
            if name in ['17','15']:
                m[app]['lh_equal']=sum(x[app] and x[app]['score'] is not None and x[app]['score']==x['case']['lh_actual'] for x in rows)
                m[app]['lh_within2']=sum(x[app] and x[app]['score'] is not None and abs(x[app]['score']-x['case']['lh_actual'])<=2 for x in rows)
        ok=[x for x in rows if x['ours'] and x['jys'] and not x['ours'].get('error') and not x['jys'].get('error')]
        m['verdict_equal']=sum(x['ours']['verdict']==x['jys']['verdict'] for x in ok)
        comparable=[x for x in ok if x['ours']['score'] is not None and x['jys']['score'] is not None]
        m['score_comparable']=len(comparable);m['score_equal']=sum(x['ours']['score']==x['jys']['score'] for x in comparable)
        m['different_parcels']=sum(set(x['ours']['pnus'])!=set(x['jys']['pnus']) for x in ok)
        metrics[name]=m
    save(OUT/'comparison.json',{'generated_at':now(),'metrics':metrics,'groups':groups})
    print(json.dumps(metrics,ensure_ascii=False,indent=2))
    print('VERDICT_DIFF',json.dumps([x for x in groups['118'] if x['ours']['verdict']!=x['jys']['verdict']],ensure_ascii=False))

def detail(args):
    data=readj(OUT/'comparison.json')
    for row in data['groups']['118']:
        a,b=row['ours'],row['jys']
        if a['pnus'][0]!=b['pnus'][0]: print('PNU_DIFFERENCE',row['key'],a['pnus'][0],b['pnus'][0])
    for key in (args.ids or '029_house_general,034_house_general,084_house_general,002_officetel_youth,077_officetel_youth,060_house_youth,095_officetel_newlywed,108_officetel_general').split(','):
        for app in ['ours','jys']:
            d=readj(OUT/'raw'/app/f'{key}.json')['result'];print(app,key)
            if app=='ours':
                print('SITE',json.dumps([{k:p.get(k) for k in ['pnu','address','geometry_source','geometry_note']} for p in d['site']['parcels']],ensure_ascii=False))
                print('REVIEW',json.dumps(d['stage_one']['review_reasons'],ensure_ascii=False))
                for ax in d['stage_two']['criteria']:
                    if ax['key'] not in ['living','education']:continue
                    print('AXIS',ax['key'],ax['awarded'])
                    for g in ax['groups']:
                        print(g['key'],json.dumps(g.get('hits',[])[:4],ensure_ascii=False))
            else:
                print('SITE',json.dumps(d['site'].get('notes'),ensure_ascii=False))
                for ax in d['stage2']:
                    if ax['axis'] not in ['LIVING','EDUCATION']:continue
                    print('AXIS',ax['axis'],ax['score'])
                    print(json.dumps([{k:f.get(k) for k in ['facility_name','facility_address','facility_pnu','distance_m_raw','measurement_type','reason_text']} for f in ax.get('used_facilities',[])],ensure_ascii=False))

def main():
    p=argparse.ArgumentParser(); p.add_argument('command'); p.add_argument('--section'); p.add_argument('--meeting',action='store_true'); p.add_argument('--all',action='store_true'); p.add_argument('--group',default='17'); p.add_argument('--app',default='both'); p.add_argument('--ids'); p.add_argument('--force',action='store_true'); args=p.parse_args()
    {'inspect':inspect,'sheets':sheets,'prepare':prepare,'run':run,'audit':audit,'control':control,'summarize':summarize,'detail':detail}[args.command](args)

if __name__=='__main__': main()
